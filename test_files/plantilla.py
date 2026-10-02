import openpyxl

# 1. Cargar el archivo original de Ingrid
wb = openpyxl.load_workbook("Estimacion 1 Normal -HGZ1-HEM OK - copia.XLSX")

# 2. Identificar y limpiar filas de datos en las hojas de trabajo
# Ajusta el nombre exacto de la hoja de cálculo de generadores/estimación si aplica
hojas_a_limpiar = wb.sheetnames

for nombre in hojas_a_limpiar:
    ws = wb[nombre]
    # Se limpian los valores a partir de la fila donde inician los conceptos
    # conservando encabezados institucionales y títulos (usualmente fila 10 o 12 en IMSS)
    fila_inicio_datos = 12
    
    for row in ws.iter_rows(min_row=fila_inicio_datos, max_row=ws.max_row):
        for cell in row:
            # Preservar fórmulas (que inician con '=') y limpiar solo valores fijos
            if cell.value and not str(cell.value).startswith('='):
                cell.value = None

# 3. Guardar la plantilla neutra lista para producción
nombre_plantilla = "plantilla_maestra_estimacion_imss.xlsx"
wb.save(nombre_plantilla)
print(f"Plantilla lista: {nombre_plantilla}")