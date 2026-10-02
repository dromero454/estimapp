import streamlit as st
from supabase import create_client, Client
import uuid
import pandas as pd
from PIL import Image
import io
import datetime
import matplotlib.pyplot as plt

# Librerías para generación de PDF corporativo
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

# Librería para inyección de Excel Institucional e imágenes
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
import requests

st.set_page_config(page_title="Estimapp", page_icon="🏗", layout="wide")

# Conexión con Supabase
url = st.secrets["SUPABASE_URL"]
key = st.secrets["SUPABASE_KEY"]
supabase: Client = create_client(url, key)

# =============================================================
# ESTILOS CSS
# =============================================================
st.markdown("""
<style>
header[data-testid="stHeader"] { display: none !important; }
.block-container { padding-top: 0.5rem !important; padding-bottom: 2rem !important; }
div[data-testid="stElementContainer"]:has(.sticky-header) {
    position: sticky !important; top: 0 !important; z-index: 99 !important;
    background-color: var(--background-color, #ffffff) !important;
}
.sticky-header {
    background-color: var(--background-color, #ffffff) !important;
    padding-top: 2px; padding-bottom: 2px;
}
div[data-baseweb="tab-list"], div[role="tablist"] {
    position: sticky !important; top: 68px !important; z-index: 98 !important;
    background-color: var(--background-color, #ffffff) !important;
    padding-top: 4px !important; padding-bottom: 4px !important;
    border-bottom: 1px solid rgba(0, 0, 0, 0.08) !important;
}
button[data-baseweb="tab"], button[data-baseweb="tab"] p, button[data-baseweb="tab"] span {
    font-size: 1.15rem !important; font-weight: 500 !important;
}
</style>

<div class="sticky-header">
    <h1 style='margin: 0; padding: 0; font-size: 2.35rem; line-height: 1.15; font-weight: 700;'>
        Estimapp <span style='font-family: "Segoe UI Emoji", "Apple Color Emoji", "Noto Color Emoji";'>🏗</span>
    </h1>
    <div style='color: #6c757d; font-size: 1.05rem; margin-top: 2px; margin-bottom: 4px;'>
        Control de avance físico y financiero de obra
    </div>
</div>
""", unsafe_allow_html=True)

# =============================================================
# FUNCIONES DE CACHÉ Y NORMALIZACIÓN
# =============================================================
@st.cache_data(show_spinner=False)
def get_proyectos():
    res = supabase.table("proyectos").select("*").order("id").execute()
    return res.data or []

@st.cache_data(show_spinner=False)
def get_biblioteca_categorias():
    res = supabase.table("biblioteca_conceptos").select("categoria").execute()
    cats = sorted(list(set(c["categoria"] for c in res.data))) if res.data else []
    if "IMSS" not in cats: cats.insert(0, "IMSS")
    if "Poder Judicial de la Federación" not in cats: cats.insert(1, "Poder Judicial de la Federación")
    return cats

@st.cache_data(show_spinner=False)
def get_biblioteca_conceptos(categoria: str):
    res = supabase.table("biblioteca_conceptos").select("*").eq("categoria", categoria).order("clave").execute()
    return res.data or []

@st.cache_data(show_spinner=False)
def get_conceptos(id_proyecto: int):
    res = supabase.table("catalogo_conceptos").select("*").eq("id_proyecto", id_proyecto).order("id").execute()
    return res.data or []

@st.cache_data(show_spinner=False)
def get_estimaciones(id_proyecto: int):
    res = supabase.table("estimaciones").select("*").eq("id_proyecto", id_proyecto).order("num_periodo").execute()
    return res.data or []

@st.cache_data(show_spinner=False)
def get_mediciones(id_estimacion: int):
    res = supabase.table("mediciones_campo").select(
        "id, localizacion, eje, tramo, largo, ancho, alto, piezas, cantidad_total, url_foto, url_croquis, id_concepto, catalogo_conceptos(clave, unidad, precio_unitario)"
    ).eq("id_estimacion", id_estimacion).order("id").execute()
    return res.data or []

def normalizar_unidad(u: str) -> str:
    if pd.isna(u) or not u or str(u).strip() == "": return "s/u"
    u_up = str(u).strip().upper()
    if u_up in ["M2", "M²"]: return "m²"
    if u_up in ["M3", "M³"]: return "m³"
    if u_up in ["LITRO", "LITROS", "LT", "LTS", "L", "ML"]: return "litros"
    if u_up in ["KG", "KILOGRAMO", "KILOS"]: return "kg"
    if u_up in ["PZA", "PZAS", "PIEZA", "PIEZAS"]: return "pza"
    if u_up in ["LOTE"]: return "lote"
    if u_up in ["TRAMO"]: return "tramo"
    if u_up in ["JGO", "JUEGO", "JGOS"]: return "jgo"
    if u_up in ["M", "METRO", "METROS"]: return "m"
    if u_up in ["H", "HR", "HORAS", "HORA", "MANO DE OBRA", "MANO DE OBRA (H)"]: return "mano de obra (h)"
    if u_up in ["M3/KM", "M³/KM", "M3 / KM", "M³ / KM"]: return "m³/km"
    return u_up.lower()

unidades_list = ["m²", "m³", "litros", "kg", "pza", "lote", "jgo", "tramo", "m", "mano de obra (h)", "m³/km", "s/u"]

def generar_plantilla_excel(tipo="biblioteca"):
    output = io.BytesIO()
    if tipo == "biblioteca": df = pd.DataFrame(columns=["Especialidad", "Clave", "Descripcion", "Unidad", "Precio_Unitario"])
    else: df = pd.DataFrame(columns=["Especialidad", "Clave", "Descripcion", "Unidad", "Cantidad_Contratada", "Precio_Unitario"])
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Plantilla')
    return output.getvalue()

def extraer_nombre_archivo(url_publica: str) -> str:
    if not url_publica: return ""
    return url_publica.split("/")[-1].split("?")[0]

def optimizar_imagen(archivo_subido, max_dim=1280, calidad=80):
    try:
        img = Image.open(archivo_subido)
        if img.mode in ("RGBA", "P"): img = img.convert("RGB")
        img.thumbnail((max_dim, max_dim), Image.Resampling.LANCZOS)
        buffer = io.BytesIO()
        img.save(buffer, format="JPEG", quality=calidad, optimize=True)
        return buffer.getvalue(), "image/jpeg"
    except Exception:
        return archivo_subido.getvalue(), archivo_subido.type

# Callbacks para categorías
def crear_categoria_callback():
    cat_nombre = st.session_state.get("input_nueva_cat", "").strip()
    if cat_nombre:
        try:
            supabase.table("biblioteca_conceptos").insert({
                "categoria": cat_nombre,
                "clave": "INIT",
                "descripcion": "Concepto inicial (Puedes borrarlo)",
                "unidad": "pza",
                "precio_referencial": 0
            }).execute()
            get_biblioteca_categorias.clear()
            st.session_state["sel_cat_admin"] = cat_nombre
        except Exception:
            pass
    st.session_state["input_nueva_cat"] = ""

def eliminar_categoria_callback():
    cat_a_borrar = st.session_state.get("sel_cat_admin")
    if cat_a_borrar:
        try:
            supabase.table("biblioteca_conceptos").delete().eq("categoria", cat_a_borrar).execute()
            get_biblioteca_categorias.clear()
            get_biblioteca_conceptos.clear()
            st.session_state["sel_cat_admin"] = "IMSS"
        except Exception:
            pass

# =============================================================
# MOTOR DE GENERACIÓN DE EXCEL INSTITUCIONAL CON OPENPYXL
# =============================================================
def descargar_plantilla_supabase(bucket_name, file_name):
    url_plantilla = supabase.storage.from_(bucket_name).get_public_url(file_name)
    response = requests.get(url_plantilla)
    if response.status_code == 200:
        return io.BytesIO(response.content)
    else:
        raise Exception(f"No se pudo descargar la plantilla de {url_plantilla}")

def set_cell_value(ws, col, row, value):
    """
    Función francotiradora: Escribe el valor en la celda indicada.
    Si la celda pertenece a un bloque combinado (MergedCell),
    busca la coordenada principal (top-left) y escribe el dato ahí para evitar errores.
    """
    coord = f"{col}{row}"
    cell = ws[coord]
    
    if type(cell).__name__ == 'MergedCell':
        for rng in ws.merged_cells.ranges:
            if cell.row >= rng.min_row and cell.row <= rng.max_row and cell.column >= rng.min_col and cell.column <= rng.max_col:
                ws.cell(row=rng.min_row, column=rng.min_col).value = value
                break
    else:
        cell.value = value

def inyectar_datos_excel_imss(plantilla_bytes, proy_info, estimacion_info, conceptos_cat, estimaciones_dash):
    wb = openpyxl.load_workbook(plantilla_bytes)
    
    # 1. Pestaña DATOS
    monto_contratado_total = sum(float(c.get("cantidad_contratada") or 0.0) * float(c.get("precio_unitario") or 0.0) for c in conceptos_cat)
    ws_datos = wb["DATOS"]
    set_cell_value(ws_datos, 'B', 2, proy_info.get('nombre_obra', ''))
    set_cell_value(ws_datos, 'B', 4, proy_info.get('descripcion_sintetica', ''))
    set_cell_value(ws_datos, 'B', 6, proy_info.get('ubicacion', ''))
    set_cell_value(ws_datos, 'B', 8, proy_info.get('unidad', ''))
    set_cell_value(ws_datos, 'B', 12, proy_info.get('contratista', ''))
    set_cell_value(ws_datos, 'B', 14, proy_info.get('residente_obra', ''))
    set_cell_value(ws_datos, 'B', 16, proy_info.get('contrato_no', ''))
    set_cell_value(ws_datos, 'B', 20, f"DEL {estimacion_info.get('periodo_inicio', '')} AL {estimacion_info.get('periodo_fin', '')}")
    set_cell_value(ws_datos, 'B', 23, datetime.datetime.now().strftime("%d/%m/%Y"))
    set_cell_value(ws_datos, 'B', 31, f"{estimacion_info.get('num_periodo', 1)} ({estimacion_info.get('estado', 'NORMAL').upper()})")
    set_cell_value(ws_datos, 'B', 35, monto_contratado_total)

    # 2. Pestaña CATALOGO
    ws_cat = wb["CATALOGO"]
    fila_cat = 14 # Ajustado para saltar los encabezados de la plantilla vacía
    for c in conceptos_cat:
        set_cell_value(ws_cat, 'A', fila_cat, c['clave'])
        set_cell_value(ws_cat, 'B', fila_cat, c['descripcion'])
        set_cell_value(ws_cat, 'C', fila_cat, normalizar_unidad(c['unidad']))
        cant = float(c.get('cantidad_contratada') or 0.0)
        pu = float(c.get('precio_unitario') or 0.0)
        set_cell_value(ws_cat, 'D', fila_cat, cant)
        set_cell_value(ws_cat, 'E', fila_cat, pu)
        set_cell_value(ws_cat, 'F', fila_cat, "") 
        set_cell_value(ws_cat, 'G', fila_cat, cant * pu)
        set_cell_value(ws_cat, 'H', fila_cat, (cant * pu / monto_contratado_total) if monto_contratado_total > 0 else 0)
        fila_cat += 1

    # Pre-cálculos cruzados para Estimacion y Generador
    acumulados_cronologicos = {c["id"]: 0.0 for c in conceptos_cat}
    conceptos_con_avance_periodo = []
    mediciones_actual = []

    for e in estimaciones_dash:
        meds_e = get_mediciones(e["id"])
        if e["id"] == estimacion_info["id"]:
            mediciones_actual = meds_e
            
        meds_por_conc_e = {}
        for m in meds_e:
            c_id = m.get("id_concepto")
            meds_por_conc_e[c_id] = meds_por_conc_e.get(c_id, 0.0) + float(m.get("cantidad_total") or 0.0)
            
        for c in conceptos_cat:
            c_id = c["id"]
            cant_periodo = meds_por_conc_e.get(c_id, 0.0)
            acumulados_cronologicos[c_id] += cant_periodo
            
            if e["id"] == estimacion_info["id"] and cant_periodo > 0:
                conceptos_con_avance_periodo.append({
                    "id": c_id,
                    "clave": c['clave'],
                    "descripcion": c['descripcion'],
                    "unidad": normalizar_unidad(c['unidad']),
                    "cant_contratada": float(c.get('cantidad_contratada') or 0.0),
                    "pu": float(c.get('precio_unitario') or 0.0),
                    "cant_periodo": cant_periodo,
                    "cant_acumulada": acumulados_cronologicos[c_id]
                })

    # 3. Pestaña Estimacion
    ws_est = wb["Estimacion"]
    fila_est = 14
    for c_idx, c in enumerate(conceptos_con_avance_periodo, start=1):
        set_cell_value(ws_est, 'A', fila_est, c_idx)
        set_cell_value(ws_est, 'B', fila_est, c['descripcion'])
        set_cell_value(ws_est, 'E', fila_est, c['unidad'])
        set_cell_value(ws_est, 'G', fila_est, c['clave'])
        set_cell_value(ws_est, 'H', fila_est, c['cant_contratada'])
        set_cell_value(ws_est, 'I', fila_est, c['cant_acumulada'])
        set_cell_value(ws_est, 'J', fila_est, c['cant_acumulada'] - c['cant_periodo'])
        set_cell_value(ws_est, 'K', fila_est, c['cant_periodo'])
        set_cell_value(ws_est, 'L', fila_est, c['pu'])
        set_cell_value(ws_est, 'M', fila_est, c['cant_periodo'] * c['pu'])
        fila_est += 1

    # 4. Pestaña Generador
    ws_gen = wb["Generador"]
    fila_gen = 14 # Alineado para caer directamente en la zona de celdas libres
    for c in conceptos_con_avance_periodo:
        set_cell_value(ws_gen, 'A', fila_gen, c['clave'])
        set_cell_value(ws_gen, 'B', fila_gen, c['descripcion'])
        set_cell_value(ws_gen, 'J', fila_gen, c['unidad'])
        # No ponemos cant_periodo total aquí porque el IMSS prefiere ver el desglose fila por fila
        fila_gen += 1
        
        meds_conc = [m for m in mediciones_actual if m['id_concepto'] == c['id']]
        for m in meds_conc:
            set_cell_value(ws_gen, 'B', fila_gen, m.get('localizacion', ''))
            set_cell_value(ws_gen, 'D', fila_gen, m.get('eje', ''))
            set_cell_value(ws_gen, 'E', fila_gen, m.get('tramo', ''))
            set_cell_value(ws_gen, 'F', fila_gen, float(m.get('ancho') or 0.0))
            set_cell_value(ws_gen, 'G', fila_gen, float(m.get('largo') or 0.0))
            set_cell_value(ws_gen, 'H', fila_gen, float(m.get('alto') or 0.0))
            set_cell_value(ws_gen, 'I', fila_gen, float(m.get('piezas') or 1.0))
            set_cell_value(ws_gen, 'J', fila_gen, c['unidad'])
            set_cell_value(ws_gen, 'K', fila_gen, float(m.get('cantidad_total') or 0.0))
            fila_gen += 1
        fila_gen += 1

    # 5. Croquis y Fotografías (Inyección binaria)
    ws_croquis = wb["Croquis"]
    ws_fotos = wb["Bit. Fotografica"]
    fila_croquis = 8
    fila_fotos = 8

    for m in mediciones_actual:
        c_ref = next((c for c in conceptos_cat if c['id'] == m['id_concepto']), {})
        c_clave = c_ref.get('clave', 'S/C')
        
        if m.get('url_croquis'):
            try:
                resp = requests.get(m['url_croquis'], timeout=10)
                if resp.status_code == 200:
                    img_stream = io.BytesIO(resp.content)
                    img = OpenpyxlImage(img_stream)
                    img.width, img.height = 400, 300
                    ws_croquis.add_image(img, f'B{fila_croquis}')
                    set_cell_value(ws_croquis, 'A', fila_croquis, f"Concepto: {c_clave} - Loc: {m.get('localizacion','')}")
                    fila_croquis += 18
            except Exception:
                pass
                
        if m.get('url_foto'):
            try:
                resp = requests.get(m['url_foto'], timeout=10)
                if resp.status_code == 200:
                    img_stream = io.BytesIO(resp.content)
                    img = OpenpyxlImage(img_stream)
                    img.width, img.height = 400, 300
                    ws_fotos.add_image(img, f'B{fila_fotos}')
                    set_cell_value(ws_fotos, 'A', fila_fotos, f"Concepto: {c_clave} - Loc: {m.get('localizacion','')}")
                    fila_fotos += 18
            except Exception:
                pass

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


# =============================================================
# MOTOR DE GENERACIÓN DE REPORTE EJECUTIVO EN PDF
# =============================================================
@st.cache_data(show_spinner=False)
def generar_pdf_resumen_ejecutivo(proy_info, monto_cont, monto_est, saldo_ejercer, pct_global, df_conceptos):
    pdf_buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        pdf_buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    story = []
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=15,
        leading=18,
        textColor=colors.HexColor("#1A2530"),
        spaceAfter=2,
        fontName="Helvetica-Bold"
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#555555"),
        spaceAfter=10
    )
    table_text = ParagraphStyle(
        'TableText',
        parent=styles['Normal'],
        fontSize=7,
        leading=9,
        textColor=colors.HexColor("#222222")
    )
    table_head = ParagraphStyle(
        'TableHead',
        parent=styles['Normal'],
        fontSize=7,
        leading=9,
        textColor=colors.white,
        fontName="Helvetica-Bold"
    )

    story.append(Paragraph("ESTIMAPP | INFORME EJECUTIVO DE CONTROL PRESUPUESTAL", title_style))
    story.append(Paragraph(f"Fecha de corte y emisión: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')} | Sistema Central de Estimaciones", subtitle_style))

    # 1. Ficha técnica del proyecto
    info_data = [
        [
            Paragraph("<b>Obra:</b>", table_text),
            Paragraph(f"{proy_info.get('nombre_obra', 'N/D')}", table_text),
            Paragraph("<b>Contrato N°:</b>", table_text),
            Paragraph(f"{proy_info.get('contrato_no', 'S/N')}", table_text)
        ],
        [
            Paragraph("<b>Ubicación:</b>", table_text),
            Paragraph(f"{proy_info.get('ubicacion', 'N/D')}", table_text),
            Paragraph("<b>Licitación:</b>", table_text),
            Paragraph(f"{proy_info.get('concurso_no', 'S/N')}", table_text)
        ],
        [
            Paragraph("<b>Contratista:</b>", table_text),
            Paragraph(f"{proy_info.get('contratista', 'N/D')}", table_text),
            Paragraph("<b>Unidad:</b>", table_text),
            Paragraph(f"{proy_info.get('unidad', 'N/D')}", table_text)
        ],
        [
            Paragraph("<b>Residente:</b>", table_text),
            Paragraph(f"{proy_info.get('residente_obra', 'N/D')}", table_text),
            Paragraph("<b>Alcance:</b>", table_text),
            Paragraph(f"{proy_info.get('descripcion_sintetica', 'General')[:50]}...", table_text)
        ]
    ]
    t_info = Table(info_data, colWidths=[65, 230, 75, 170])
    t_info.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#F8FAFC")),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#CBD5E1")),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#E2E8F0")),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
    ]))
    story.append(t_info)
    story.append(Spacer(1, 8))

    # 2. Métricas Financieras (KPIs)
    kpi_data = [
        [
            Paragraph(f"<para align=center><b>Monto Contratado Total</b><br/><font size=10 color='#1A2530'><b>${monto_cont:,.2f}</b></font></para>", table_text),
            Paragraph(f"<para align=center><b>Monto Estimado Acumulado</b><br/><font size=10 color='#00875A'><b>${monto_est:,.2f}</b></font></para>", table_text),
            Paragraph(f"<para align=center><b>Saldo por Ejercer</b><br/><font size=10 color='#D32F2F'><b>${saldo_ejercer:,.2f}</b></font></para>", table_text),
            Paragraph(f"<para align=center><b>Avance Global</b><br/><font size=10 color='#0288D1'><b>{pct_global:.2f}%</b></font></para>", table_text)
        ]
    ]
    t_kpi = Table(kpi_data, colWidths=[135, 135, 135, 135])
    t_kpi.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.white),
        ('BOX', (0,0), (-1,-1), 1.2, colors.HexColor("#00875A")),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#E2E8F0")),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
    ]))
    story.append(t_kpi)
    story.append(Spacer(1, 8))

    # 3. Gráfica en memoria
    fig, ax = plt.subplots(figsize=(7.2, 1.0), dpi=150)
    cats = ['Balance Contractual']
    ax.barh(cats, [monto_est], color='#00875A', height=0.45, label=f'Ejercido (${monto_est:,.2f} - {pct_global:.1f}%)')
    ax.barh(cats, [saldo_ejercer], left=[monto_est], color='#CBD5E1', height=0.45, label=f'Saldo (${saldo_ejercer:,.2f})')
    limite = max(monto_cont * 1.05, 1.0)
    ax.set_xlim(0, limite)
    ax.set_xlabel('Monto en MXN', fontsize=7)
    ax.legend(loc='lower center', bbox_to_anchor=(0.5, 1.01), ncol=2, frameon=False, fontsize=7)
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.spines['left'].set_visible(False)
    ax.tick_params(axis='both', which='both', labelsize=7, left=False)
    plt.tight_layout()

    chart_buf = io.BytesIO()
    plt.savefig(chart_buf, format='png', bbox_inches='tight')
    plt.close(fig)
    chart_buf.seek(0)
    story.append(RLImage(chart_buf, width=540, height=75))
    story.append(Spacer(1, 8))

    # 4. Tabla consolidada de conceptos
    table_rows = [
        [
            Paragraph("<para align=center>#</para>", table_head),
            Paragraph("<para align=center>Clave</para>", table_head),
            Paragraph("<para align=center>Und</para>", table_head),
            Paragraph("<para align=center>Cant. Cont.</para>", table_head),
            Paragraph("<para align=center>Cant. Est.</para>", table_head),
            Paragraph("<para align=center>P.U. ($)</para>", table_head),
            Paragraph("<para align=center>Importe Est. ($)</para>", table_head),
            Paragraph("<para align=center>Saldo Fin. ($)</para>", table_head),
            Paragraph("<para align=center>% Avance</para>", table_head),
        ]
    ]

    for idx, row in df_conceptos.iterrows():
        p_avance = float(row.get("% Avance", 0.0))
        color_pct = "#D32F2F" if p_avance > 100.0 else "#00875A"
        table_rows.append([
            Paragraph(f"<para align=center>{idx+1}</para>", table_text),
            Paragraph(f"{row.get('Clave', '')}", table_text),
            Paragraph(f"<para align=center>{row.get('Unidad', '')}</para>", table_text),
            Paragraph(f"<para align=right>{float(row.get('Cant. Contratada', 0)):,.2f}</para>", table_text),
            Paragraph(f"<para align=right>{float(row.get('Cant. Acumulada', 0)):,.2f}</para>", table_text),
            Paragraph(f"<para align=right>${float(row.get('P.U. ($)', 0)):,.2f}</para>", table_text),
            Paragraph(f"<para align=right>${float(row.get('Importe Acumulado ($)', 0)):,.2f}</para>", table_text),
            Paragraph(f"<para align=right>${float(row.get('Saldo Financiero ($)', 0)):,.2f}</para>", table_text),
            Paragraph(f"<para align=center><font color='{color_pct}'><b>{p_avance:.1f}%</b></font></para>", table_text),
        ])

    t_table = Table(table_rows, colWidths=[23, 72, 35, 65, 65, 60, 75, 75, 70], repeatRows=1)
    t_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#1A2530")),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#E2E8F0")),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#94A3B8")),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 2.5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 2.5),
    ]))
    story.append(t_table)
    story.append(Spacer(1, 15))

    # 5. Cuadro oficial de firmas
    residente_nom = proy_info.get('residente_obra') or 'Supervisión de Obra'
    contratista_nom = proy_info.get('contratista') or 'Contratista'
    sig_data = [
        [
            Paragraph(f"<para align=center>____________________________________<br/><b>{residente_nom}</b><br/>Residente de Obra / Supervisión</para>", table_text),
            Paragraph(f"<para align=center>____________________________________<br/><b>{contratista_nom}</b><br/>Empresa Contratista</para>", table_text)
        ]
    ]
    t_sig = Table(sig_data, colWidths=[270, 270])
    t_sig.setStyle(TableStyle([
        ('TOPPADDING', (0,0), (-1,-1), 15),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
    ]))
    story.append(KeepTogether(t_sig))

    doc.build(story)
    return pdf_buffer.getvalue()

keys = ["del_proy_counter", "del_conc_counter", "del_est_counter", "del_med_counter", "cap_counter", "dim_counter", "bib_del_counter", "up_bib_key", "up_proy_key", "ms_lote_key", "new_cat_key"]
for k in keys:
    if k not in st.session_state: st.session_state[k] = 0

# -------------------------------------------------------------
# DEFINICIÓN DE PESTAÑAS
# -------------------------------------------------------------
tab_dashboard, tab_captura, tab_estimaciones, tab_catalogo, tab_biblioteca, tab_proyectos = st.tabs([
    "📊 Control Presupuestal", "📐 Captura en Campo", "📑 Estimaciones", 
    "📚 Catálogo del Proyecto", "📖 Biblioteca Maestra de Conceptos", "🏢 Proyectos"
])

lista_proyectos = get_proyectos()
proyectos_dict = {p["nombre_obra"]: p["id"] for p in lista_proyectos} if lista_proyectos else {}

# -------------------------------------------------------------
# TAB 1: CONTROL PRESUPUESTAL (DASHBOARD + REPORTE PDF EJECUTIVO)
# -------------------------------------------------------------
with tab_dashboard:
    st.subheader("Control Presupuestal y Balance Contractual")
    if not proyectos_dict: st.info("Registra un proyecto para visualizar el análisis financiero.")
    else:
        proy_sel_dash = st.selectbox("Proyecto a Auditar", list(proyectos_dict.keys()), key="dash_proy")
        id_proy_dash = proyectos_dict[proy_sel_dash]
        proy_obj_actual = next((p for p in lista_proyectos if p["id"] == id_proy_dash), {})

        conceptos_dash = get_conceptos(id_proy_dash)
        estimaciones_dash = get_estimaciones(id_proy_dash)

        if not conceptos_dash: st.warning("Este proyecto no tiene conceptos en su catálogo.")
        else:
            dict_conceptos_info = {c["id"]: c for c in conceptos_dash}
            monto_contratado_total = sum(float(c.get("cantidad_contratada") or 0.0) * float(c.get("precio_unitario") or 0.0) for c in conceptos_dash)
            meds_por_est_conc = {}
            monto_estimado_global = 0.0
            conceptos_con_estimacion = set()

            for e in estimaciones_dash:
                meds = get_mediciones(e["id"])
                for m in meds:
                    c_id = m.get("id_concepto")
                    if c_id in dict_conceptos_info:
                        cant = float(m["cantidad_total"] or 0.0)
                        pu = float(dict_conceptos_info[c_id].get("precio_unitario") or 0.0)
                        meds_por_est_conc[(e["id"], c_id)] = meds_por_est_conc.get((e["id"], c_id), 0.0) + cant
                        monto_estimado_global += (cant * pu)
                        conceptos_con_estimacion.add(c_id)

            saldo_por_ejercer = round(monto_contratado_total - monto_estimado_global, 2)
            pct_global = round((monto_estimado_global / monto_contratado_total * 100), 2) if monto_contratado_total > 0 else 0.0

            kpi1, kpi2, kpi3, kpi4 = st.columns(4)
            kpi1.metric("Monto Contratado Total", f"${monto_contratado_total:,.2f}")
            kpi2.metric("Monto Estimado Acumulado", f"${monto_estimado_global:,.2f}")
            kpi3.metric("Saldo por Ejercer (Meta = $0)", f"${saldo_por_ejercer:,.2f}", delta=f"{saldo_por_ejercer:,.2f}", delta_color="inverse")
            kpi4.metric("Avance Financiero Global", f"{pct_global:.2f}%")

            st.progress(min(pct_global / 100.0, 1.0))
            st.markdown("---")

            filas_dash = []
            acumulados_cronologicos = {c["id"]: 0.0 for c in conceptos_dash}

            for e in estimaciones_dash:
                num_p = e["num_periodo"]
                label_est = f"Estimación #{num_p}"
                for c in conceptos_dash:
                    c_id = c["id"]
                    if (e["id"], c_id) in meds_por_est_conc:
                        cant_periodo = round(meds_por_est_conc[(e["id"], c_id)], 3)
                        acumulados_cronologicos[c_id] += cant_periodo
                        cant_acum = round(acumulados_cronologicos[c_id], 3)
                        cant_cont = float(c.get("cantidad_contratada") or 0.0)
                        pu = float(c.get("precio_unitario") or 0.0)
                        imp_periodo, imp_acum = round(cant_periodo * pu, 2), round(cant_acum * pu, 2)
                        saldo_vol, saldo_fin = round(cant_cont - cant_acum, 3), round((cant_cont * pu) - imp_acum, 2)
                        avance = round((cant_acum / cant_cont * 100), 1) if cant_cont > 0 else 0.0

                        if avance > 100.0: estatus_badge = f"🚨 Sobregirado (+{round(avance - 100.0, 1)}%)"
                        elif avance == 100.0: estatus_badge = "✅ Concluido (100%)"
                        else: estatus_badge = "🟢 En proceso"

                        filas_dash.append({
                            "Estimación": label_est, "Periodo": f"{e['periodo_inicio']} al {e['periodo_fin']}", "Estado": e.get("estado", "borrador"),
                            "Clave": c["clave"], "Unidad": normalizar_unidad(c["unidad"]),
                            "Cant. Contratada": cant_cont, "Cant. en Periodo": cant_periodo, "Cant. Acumulada": cant_acum,
                            "Saldo Físico": saldo_vol, "P.U. ($)": pu, "Importe Periodo ($)": imp_periodo,
                            "Importe Acumulado ($)": imp_acum, "Saldo Financiero ($)": saldo_fin,
                            "% Avance": avance, "Estatus": estatus_badge, "_sort_est": num_p, "_sort_avance": avance
                        })

            for c in conceptos_dash:
                if c["id"] not in conceptos_con_estimacion:
                    cant_cont, pu = float(c.get("cantidad_contratada") or 0.0), float(c.get("precio_unitario") or 0.0)
                    filas_dash.append({
                        "Estimación": "Sin estimar", "Periodo": "—", "Estado": "pendiente",
                        "Clave": c["clave"], "Unidad": normalizar_unidad(c["unidad"]),
                        "Cant. Contratada": cant_cont, "Cant. en Periodo": 0.0, "Cant. Acumulada": 0.0,
                        "Saldo Físico": cant_cont, "P.U. ($)": pu, "Importe Periodo ($)": 0.0,
                        "Importe Acumulado ($)": 0.0, "Saldo Financiero ($)": round(cant_cont * pu, 2),
                        "% Avance": 0.0, "Estatus": "⚪ Sin iniciar", "_sort_est": 9999, "_sort_avance": 0.0
                    })

            if filas_dash:
                df_dash = pd.DataFrame(filas_dash)
                df_dash = df_dash.sort_values(by=["_sort_est", "_sort_avance"], ascending=[True, False]).reset_index(drop=True)
                df_dash.insert(3, "#", range(1, len(df_dash) + 1))
                df_dash = df_dash.drop(columns=["_sort_est", "_sort_avance"])

                # Botón exclusivo de descarga de Resumen Ejecutivo en PDF
                col_titulo_rep, col_btn_pdf = st.columns([3, 1])
                with col_titulo_rep:
                    st.markdown("##### Desglose por Estimación y Concepto (Auditoría Integral de Avance):")
                with col_btn_pdf:
                    pdf_bytes = generar_pdf_resumen_ejecutivo(
                        proy_info=proy_obj_actual,
                        monto_cont=monto_contratado_total,
                        monto_est=monto_estimado_global,
                        saldo_ejercer=saldo_por_ejercer,
                        pct_global=pct_global,
                        df_conceptos=df_dash
                    )
                    st.download_button(
                        label="📄 Exportar Resumen Ejecutivo (PDF)",
                        data=pdf_bytes,
                        file_name=f"Resumen_Ejecutivo_{proy_obj_actual.get('contrato_no', 'Obra')}.pdf",
                        mime="application/pdf",
                        use_container_width=True
                    )

                st.dataframe(
                    df_dash,
                    column_config={
                        "Estimación": st.column_config.TextColumn("Estimación"), "Periodo": st.column_config.TextColumn("Fechas de Corte"), "Estado": st.column_config.TextColumn("Estado"),
                        "#": st.column_config.NumberColumn("#", width="small"), "Clave": st.column_config.TextColumn("Clave"), "Unidad": st.column_config.TextColumn("Unidad", width="small"),
                        "P.U. ($)": st.column_config.NumberColumn("P.U. ($)", format="$%.2f"), "Importe Periodo ($)": st.column_config.NumberColumn("Importe Periodo ($)", format="$%.2f"),
                        "Importe Acumulado ($)": st.column_config.NumberColumn("Importe Acumulado ($)", format="$%.2f"), "Saldo Financiero ($)": st.column_config.NumberColumn("Saldo Financiero ($)", format="$%.2f"),
                        "% Avance": st.column_config.ProgressColumn("% Avance", min_value=0, max_value=100, format="%.1f%%"),
                    },
                    use_container_width=True, hide_index=True
                )

# -------------------------------------------------------------
# TAB 2: CAPTURA EN CAMPO
# -------------------------------------------------------------
with tab_captura:
    st.subheader("Captura de Mediciones y Evidencia")
    if not proyectos_dict: st.warning("Configura tu proyecto primero.")
    else:
        proy_sel_cap = st.selectbox("Proyecto", list(proyectos_dict.keys()), key="cap_proy")
        id_proy_cap = proyectos_dict[proy_sel_cap]
        lista_est_cap = get_estimaciones(id_proy_cap)
        estimaciones_dict = {f"Estimación #{e['num_periodo']} ({e['estado']})": e["id"] for e in lista_est_cap} if lista_est_cap else {}
        lista_conc_cap = get_conceptos(id_proy_cap)
        conceptos_dict = {}
        if lista_conc_cap:
            for c in lista_conc_cap:
                clave_c, desc_c, u_norm = c["clave"], c["descripcion"], normalizar_unidad(c["unidad"])
                desc_limpia = desc_c[len(clave_c):].lstrip(" .:-") if desc_c.lower().startswith(clave_c.lower()) else desc_c
                conceptos_dict[f"{clave_c} — {desc_limpia[:55]}... ({u_norm})"] = c

        if not estimaciones_dict: st.error("No hay periodos de estimación abiertos para este proyecto.")
        elif not conceptos_dict: st.error("No hay conceptos en el catálogo de este proyecto.")
        else:
            c_est, c_con = st.columns([1, 2])
            est_sel = c_est.selectbox("Periodo Activo", list(estimaciones_dict.keys()))
            conc_sel = c_con.selectbox("Concepto a Cuantificar", list(conceptos_dict.keys()))
            id_est = estimaciones_dict[est_sel]
            obj_conc = conceptos_dict[conc_sel]
            id_conc, u_base, pu_conc = obj_conc["id"], normalizar_unidad(obj_conc["unidad"]), float(obj_conc.get("precio_unitario") or 0.0)

            if "last_unidad_cap" not in st.session_state: st.session_state.last_unidad_cap = u_base
            if st.session_state.last_unidad_cap != u_base:
                st.session_state.last_unidad_cap = u_base
                st.session_state.dim_counter += 1
                st.rerun()

            st.markdown("---")
            st.markdown(f"**Concepto:** `{obj_conc['clave']}` | **Unidad:** `{u_base}` | **P.U. Contratado:** `${pu_conc:,.2f}`")
            
            c_ver = st.session_state.cap_counter
            d_ver = f"{st.session_state.cap_counter}_{st.session_state.dim_counter}"

            col_loc1, col_loc2, col_loc3 = st.columns(3)
            localizacion = col_loc1.text_input("Localización / Elemento *", placeholder="Ej: VESTIDORES - MURO FONDO...", key=f"input_loc_{c_ver}")
            eje = col_loc2.text_input("Eje", placeholder="Ej: 2, A-B...", key=f"input_eje_{c_ver}")
            tramo = col_loc3.text_input("Tramo", placeholder="Ej: 1-2, EJE C...", key=f"input_tramo_{c_ver}")

            es_m3, es_m2, es_kg, es_lt = (u_base=="m³"), (u_base=="m²"), (u_base=="kg"), (u_base=="litros")
            es_h, es_m3km, es_lin = (u_base=="mano de obra (h)"), (u_base=="m³/km"), (u_base in ["m", "tramo"])
            
            tipo_geom = "Rectangular / Cuadrada"
            if es_m2: tipo_geom = st.selectbox("Tipo de Geometría:", ["Rectangular / Cuadrada", "Triangular", "Trapecio Regular"], key=f"geom_sel_{d_ver}")

            if es_m2:
                if tipo_geom == "Triangular": label_largo, label_ancho, label_alto, des_largo, des_ancho, des_alto = "Base (m)", "Altura (m)", "Alto (m)", False, False, True
                elif tipo_geom == "Trapecio Regular": label_largo, label_ancho, label_alto, des_largo, des_ancho, des_alto = "Base Mayor - B (m)", "Base Menor - b (m)", "Altura - h (m)", False, False, False
                else: label_largo, label_ancho, label_alto, des_largo, des_ancho, des_alto = "Largo (m)", "Ancho / Altura (m)", "Alto (m)", False, False, True
                des_kg, des_litros, des_horas = True, True, True
            elif es_m3km:
                label_largo, label_ancho, label_alto, des_largo, des_ancho, des_alto, des_kg, des_litros, des_horas = "Distancia (km)", "Volumen (m³)", "Alto (m)", False, False, True, True, True, True
            elif es_m3: label_largo, label_ancho, label_alto, des_largo, des_ancho, des_alto, des_kg, des_litros, des_horas = "Largo (m)", "Ancho (m)", "Alto (m)", False, False, False, True, True, True
            elif es_kg: label_largo, label_ancho, label_alto, des_largo, des_ancho, des_alto, des_kg, des_litros, des_horas = "Largo (m)", "Ancho (m)", "Alto (m)", False, True, True, False, True, True
            elif es_lt: label_largo, label_ancho, label_alto, des_largo, des_ancho, des_alto, des_kg, des_litros, des_horas = "Largo (m)", "Ancho (m)", "Alto (m)", True, True, True, True, False, True
            elif es_h: label_largo, label_ancho, label_alto, des_largo, des_ancho, des_alto, des_kg, des_litros, des_horas = "Largo (m)", "Ancho (m)", "Alto (m)", True, True, True, True, True, False
            elif es_lin: label_largo, label_ancho, label_alto, des_largo, des_ancho, des_alto, des_kg, des_litros, des_horas = "Largo (m)", "Ancho (m)", "Alto (m)", False, True, True, True, True, True
            else: label_largo, label_ancho, label_alto, des_largo, des_ancho, des_alto, des_kg, des_litros, des_horas = "Largo (m)", "Ancho (m)", "Alto (m)", True, True, True, True, True, True

            col_d1, col_d2, col_d3, col_d4, col_d5, col_d6, col_d7 = st.columns(7)
            largo = col_d1.number_input(label_largo, min_value=0.0, value=0.0, step=0.5, key=f"nl_{d_ver}", disabled=des_largo)
            ancho = col_d2.number_input(label_ancho, min_value=0.0, value=0.0, step=0.5, key=f"nan_{d_ver}", disabled=des_ancho)
            alto = col_d3.number_input(label_alto, min_value=0.0, value=0.0, step=0.5, key=f"nal_{d_ver}", disabled=des_alto)
            kilos = col_d4.number_input("Kilos (kg)", min_value=0.0, value=0.0, step=0.5, key=f"nkg_{d_ver}", disabled=des_kg)
            litros = col_d5.number_input("Litros", min_value=0.0, value=0.0, step=0.5, key=f"nlt_{d_ver}", disabled=des_litros)
            horas = col_d6.number_input("Horas (h)", min_value=0.0, value=0.0, step=0.5, key=f"nhr_{d_ver}", disabled=des_horas)
            piezas = col_d7.number_input("Piezas", min_value=1.0, value=1.0, step=1.0, key=f"npz_{d_ver}")

            v_l, v_an, v_al = (0.0 if des_largo else float(largo)), (0.0 if des_ancho else float(ancho)), (0.0 if des_alto else float(alto))
            v_kg, v_lt, v_h = (0.0 if des_kg else float(kilos)), (0.0 if des_litros else float(litros)), (0.0 if des_horas else float(horas))

            if es_m2:
                if tipo_geom == "Triangular": calc_preview = (v_l * v_an / 2.0) * piezas
                elif tipo_geom == "Trapecio Regular": calc_preview = ((v_l + v_an) / 2.0) * v_al * piezas
                else: calc_preview = (v_l * v_an * piezas) if (v_l > 0 and v_an > 0) else 0.0
            elif es_m3km: calc_preview = v_l * v_an * piezas
            elif es_m3: calc_preview = v_l * v_an * v_al * piezas
            elif es_kg: calc_preview = (v_l * v_kg * piezas) if v_l > 0 else (v_kg * piezas)
            elif es_lt: calc_preview = v_lt * piezas
            elif es_h: calc_preview = v_h * piezas
            elif es_lin: calc_preview = v_l * piezas
            else: calc_preview = piezas

            calc_preview = round(float(calc_preview), 3)
            importe_preview = round(calc_preview * pu_conc, 2)
            st.info(f"📐 Cantidad Calculada: **{calc_preview:.3f} {u_base}** | Importe Estimado: **${importe_preview:,.2f} MXN**")

            col_f1, col_f2 = st.columns(2)
            foto = col_f1.file_uploader("Fotografía de Evidencia", type=["jpg", "jpeg", "png"], key=f"file_foto_{c_ver}")
            croquis = col_f2.file_uploader("Croquis / Plano", type=["jpg", "jpeg", "png"], key=f"file_croquis_{c_ver}")

            if st.button("💾 Guardar Medición en Generador", type="primary"):
                if calc_preview <= 0: st.warning("La cantidad debe ser mayor a 0.")
                elif not localizacion.strip(): st.warning("Debes indicar la Localización / Elemento.")
                else:
                    url_foto, url_croquis = None, None
                    if foto:
                        contenido_f, mime_f = optimizar_imagen(foto)
                        fname_f = f"{uuid.uuid4()}.jpg"
                        supabase.storage.from_("evidencias").upload(fname_f, contenido_f, {"content-type": mime_f})
                        url_foto = supabase.storage.from_("evidencias").get_public_url(fname_f)
                    if croquis:
                        contenido_c, mime_c = optimizar_imagen(croquis)
                        fname_c = f"{uuid.uuid4()}.jpg"
                        supabase.storage.from_("evidencias").upload(fname_c, contenido_c, {"content-type": mime_c})
                        url_croquis = supabase.storage.from_("evidencias").get_public_url(fname_c)

                    if es_kg: v_an = v_kg
                    elif es_lt: v_l = v_lt
                    elif es_h: v_l = v_h

                    supabase.table("mediciones_campo").insert({
                        "id_estimacion": id_est, "id_concepto": id_conc, "localizacion": localizacion.strip(),
                        "eje": eje.strip(), "tramo": tramo.strip(), "largo": v_l, "ancho": v_an, "alto": v_al,
                        "piezas": float(piezas), "cantidad_total": calc_preview, "url_foto": url_foto, "url_croquis": url_croquis
                    }).execute()
                    get_mediciones.clear()
                    st.session_state.cap_counter += 1
                    st.session_state.dim_counter += 1
                    st.success("✅ Medición guardada.")
                    st.rerun()

            st.markdown("### Mediciones registradas en este periodo:")
            st.caption("💡 *Haz doble clic sobre cualquier celda permitida para editar directamente su valor.*")
            mediciones_periodo = get_mediciones(id_est)

            if mediciones_periodo:
                rows, meds_borrar_dict, total_periodo_acum = [], {}, 0.0
                for idx, m in enumerate(mediciones_periodo, start=1):
                    clave_c, u_c, pu_c = m["catalogo_conceptos"]["clave"], normalizar_unidad(m["catalogo_conceptos"]["unidad"]), float(m["catalogo_conceptos"].get("precio_unitario") or 0.0)
                    cant_m = float(m["cantidad_total"] or 0.0)
                    imp_m = round(cant_m * pu_c, 2)
                    total_periodo_acum += imp_m
                    rows.append({
                        "#": idx, "id": m["id"], "Clave": clave_c, "Localización": m["localizacion"] or "",
                        "Eje": m["eje"] or "", "Tramo": m["tramo"] or "", "Largo / Factor": float(m["largo"] or 0.0),
                        "Ancho / Kilos": float(m["ancho"] or 0.0), "Alto": float(m["alto"] or 0.0), "Pzas": float(m["piezas"] or 1.0),
                        "Cantidad": cant_m, "Unidad": u_c, "P.U. ($)": pu_c, "Importe ($)": imp_m,
                        "Tiene Foto": "Sí" if m["url_foto"] else "No", "Tiene Croquis": "Sí" if m["url_croquis"] else "No"
                    })
                    meds_borrar_dict[f"#{idx} — {clave_c} ({m['localizacion']} — {cant_m} {u_c})"] = m

                df_meds = pd.DataFrame(rows)
                edited_meds = st.data_editor(
                    df_meds.drop(columns=["id"]),
                    column_config={
                        "#": st.column_config.NumberColumn("#", disabled=True), "Clave": st.column_config.TextColumn("Clave", disabled=True),
                        "Cantidad": st.column_config.NumberColumn("Cantidad", disabled=True), "Unidad": st.column_config.TextColumn("Unidad", disabled=True),
                        "P.U. ($)": st.column_config.NumberColumn("P.U. ($)", format="$%.2f", disabled=True), "Importe ($)": st.column_config.NumberColumn("Importe ($)", format="$%.2f", disabled=True),
                        "Tiene Foto": st.column_config.TextColumn("Foto", disabled=True), "Tiene Croquis": st.column_config.TextColumn("Croquis", disabled=True),
                    },
                    use_container_width=True, hide_index=True, key="editor_mediciones"
                )

                if "editor_mediciones" in st.session_state and st.session_state["editor_mediciones"].get("edited_rows", {}):
                    for row_str, col_vals in st.session_state["editor_mediciones"]["edited_rows"].items():
                        r_idx, up_payload = int(row_str), {}
                        id_med_up = rows[r_idx]["id"]
                        for k in ["Localización", "Eje", "Tramo"]:
                            if k in col_vals: up_payload[k.lower()] = col_vals[k].strip()
                        for k in ["Largo / Factor", "Ancho / Kilos", "Alto", "Pzas"]:
                            if k in col_vals: 
                                if k == "Largo / Factor": up_payload["largo"] = float(col_vals[k])
                                elif k == "Ancho / Kilos": up_payload["ancho"] = float(col_vals[k])
                                elif k == "Pzas": up_payload["piezas"] = float(col_vals[k])
                                else: up_payload["alto"] = float(col_vals[k])

                        if any(k in col_vals for k in ["Largo / Factor", "Ancho / Kilos", "Alto", "Pzas"]):
                            l_val = float(col_vals.get("Largo / Factor", rows[r_idx]["Largo / Factor"]))
                            an_val = float(col_vals.get("Ancho / Kilos", rows[r_idx]["Ancho / Kilos"]))
                            al_val = float(col_vals.get("Alto", rows[r_idx]["Alto"]))
                            pz_val = float(col_vals.get("Pzas", rows[r_idx]["Pzas"]))
                            u_item = rows[r_idx]["Unidad"]
                            
                            if u_item == "m³": c_calc = l_val * an_val * al_val * pz_val
                            elif u_item == "m²": c_calc = (l_val * an_val * pz_val) if (l_val > 0 and an_val > 0) else 0.0
                            elif u_item == "m³/km": c_calc = l_val * an_val * pz_val
                            elif u_item == "kg": c_calc = (l_val * an_val * pz_val) if l_val > 0 else (an_val * pz_val)
                            elif u_item in ["m", "tramo", "litros", "mano de obra (h)"]: c_calc = l_val * pz_val
                            else: c_calc = pz_val
                            up_payload["cantidad_total"] = round(c_calc, 3)

                        supabase.table("mediciones_campo").update(up_payload).eq("id", id_med_up).execute()
                        get_mediciones.clear()
                        st.toast("✅ Medición actualizada.")
                        st.rerun()

                st.metric("Total Estimado en el Periodo (Sin I.V.A.)", f"${total_periodo_acum:,.2f} MXN")

                with st.expander("🗑️️ Eliminar Mediciones (Borrado Masivo)"):
                    meds_a_borrar_labels = st.multiselect("Seleccionar mediciones a remover:", list(meds_borrar_dict.keys()), key=f"sel_med_del_{st.session_state.del_med_counter}")
                    borrar_todas_meds = st.checkbox("⚠️ Selecciona para eliminar todas las mediciones mostradas en la tabla.", key=f"chk_todas_meds_{st.session_state.del_med_counter}")
                    if borrar_todas_meds: meds_a_borrar_labels = list(meds_borrar_dict.keys())

                    if st.button("Eliminar Seleccionadas", type="primary", disabled=len(meds_a_borrar_labels)==0, key="btn_del_meds_bulk"):
                        ids_to_delete, archivos_a_borrar = [], []
                        for label in meds_a_borrar_labels:
                            obj_med_borrar = meds_borrar_dict[label]
                            ids_to_delete.append(obj_med_borrar["id"])
                            if obj_med_borrar.get("url_foto"):
                                nom_f = extraer_nombre_archivo(obj_med_borrar["url_foto"])
                                if nom_f: archivos_a_borrar.append(nom_f)
                            if obj_med_borrar.get("url_croquis"):
                                nom_c = extraer_nombre_archivo(obj_med_borrar["url_croquis"])
                                if nom_c: archivos_a_borrar.append(nom_c)

                        if archivos_a_borrar:
                            for chunk in [archivos_a_borrar[i:i+50] for i in range(0, len(archivos_a_borrar), 50)]:
                                try: supabase.storage.from_("evidencias").remove(chunk)
                                except: pass
                        if ids_to_delete:
                            for chunk in [ids_to_delete[i:i+50] for i in range(0, len(ids_to_delete), 50)]:
                                supabase.table("mediciones_campo").delete().in_("id", chunk).execute()
                        get_mediciones.clear()
                        st.session_state.del_med_counter += 1
                        st.success(f"{len(ids_to_delete)} mediciones eliminadas.")
                        st.rerun()

# -------------------------------------------------------------
# TAB 3: ESTIMACIONES
# -------------------------------------------------------------
with tab_estimaciones:
    st.subheader("Periodos de Estimación")
    if not proyectos_dict: st.warning("Registra un proyecto primero.")
    else:
        proy_sel_est = st.selectbox("Seleccionar Proyecto", list(proyectos_dict.keys()), key="est_proy")
        id_proy_est = proyectos_dict[proy_sel_est]

        col_est_nueva, col_est_baja = st.columns(2)
        with col_est_nueva:
            with st.expander("➕ Aperturar nueva estimación"):
                with st.form("form_nueva_estimacion", clear_on_submit=True):
                    col_e1, col_e2, col_e3 = st.columns(3)
                    num_periodo = col_e1.number_input("N° Estimación", min_value=1, step=1, value=1)
                    f_ini = col_e2.date_input("Fecha Inicio")
                    f_fin = col_e3.date_input("Fecha Fin")
                    if st.form_submit_button("Abrir Periodo"):
                        try:
                            supabase.table("estimaciones").insert({
                                "id_proyecto": id_proy_est, "num_periodo": int(num_periodo),
                                "periodo_inicio": str(f_ini), "periodo_fin": str(f_fin), "estado": "borrador"
                            }).execute()
                            get_estimaciones.clear()
                            st.success(f"Estimación #{num_periodo} aperturada.")
                            st.rerun()
                        except Exception as e: st.error(f"Error al aperturar: {e}")

        estimaciones_proyecto = get_estimaciones(id_proy_est)
        dict_est_borrar = {f"#{idx} — Estimación #{e['num_periodo']} ({e['periodo_inicio']} al {e['periodo_fin']})": e["id"] for idx, e in enumerate(estimaciones_proyecto, start=1)} if estimaciones_proyecto else {}

        with col_est_baja:
            with st.expander("🗑️ Eliminar Estimación"):
                if not dict_est_borrar: st.info("No hay estimaciones registradas.")
                else:
                    est_del_sel = st.selectbox("Seleccionar estimación:", list(dict_est_borrar.keys()), key=f"del_est_sel_{st.session_state.del_est_counter}")
                    st.warning("⚠️ Al eliminar la estimación también se borrarán todas sus mediciones.")
                    if st.button("Eliminar Estimación", type="primary", disabled=not st.checkbox("Confirmo eliminar esta estimación", key=f"chk_del_est_{st.session_state.del_est_counter}")):
                        supabase.table("estimaciones").delete().eq("id", dict_est_borrar[est_del_sel]).execute()
                        get_estimaciones.clear()
                        get_mediciones.clear()
                        st.session_state.del_est_counter += 1
                        st.success("Estimación eliminada.")
                        st.rerun()

        if estimaciones_proyecto:
            st.markdown("##### Listado de Estimaciones:")
            st.caption("💡 *Haz doble clic en cualquier celda para cambiar de estado o corregir las fechas de corte.*")
            df_e = pd.DataFrame(estimaciones_proyecto)
            df_e.insert(0, "#", range(1, len(df_e) + 1))
            df_editor_data = df_e[["#", "num_periodo", "periodo_inicio", "periodo_fin", "estado"]].copy()
            df_editor_data["num_periodo"] = pd.to_numeric(df_editor_data["num_periodo"], errors="coerce").fillna(1).astype(int)
            df_editor_data["periodo_inicio"] = pd.to_datetime(df_editor_data["periodo_inicio"], errors="coerce").dt.date
            df_editor_data["periodo_fin"] = pd.to_datetime(df_editor_data["periodo_fin"], errors="coerce").dt.date
            df_editor_data["estado"] = df_editor_data["estado"].fillna("borrador").astype(str)

            edited_table = st.data_editor(
                df_editor_data,
                column_config={
                    "estado": st.column_config.SelectboxColumn("Estado de la Estimación", width="medium", options=["borrador", "en_revision", "aprobada"], required=True),
                    "#": st.column_config.NumberColumn("#", disabled=True), "num_periodo": st.column_config.NumberColumn("N° Periodo", disabled=True),
                    "periodo_inicio": st.column_config.DateColumn("Fecha Inicio", format="YYYY-MM-DD"), "periodo_fin": st.column_config.DateColumn("Fecha Fin", format="YYYY-MM-DD"),
                },
                hide_index=True, use_container_width=True, key="editor_estimaciones"
            )

            if "editor_estimaciones" in st.session_state and st.session_state["editor_estimaciones"].get("edited_rows", {}):
                for row_idx_str, col_mod in st.session_state["editor_estimaciones"]["edited_rows"].items():
                    id_est_mod, up_est = estimaciones_proyecto[int(row_idx_str)]["id"], {}
                    if "estado" in col_mod: up_est["estado"] = col_mod["estado"]
                    if "periodo_inicio" in col_mod: up_est["periodo_inicio"] = str(col_mod["periodo_inicio"])
                    if "periodo_fin" in col_mod: up_est["periodo_fin"] = str(col_mod["periodo_fin"])
                    supabase.table("estimaciones").update(up_est).eq("id", id_est_mod).execute()
                    get_estimaciones.clear()
                    st.toast("✅ Estimación actualizada.")
                    st.rerun()

            # --- MODULO DE EXPORTACIÓN OFICIAL EXCEL (INYECCIÓN PURA) ---
            st.markdown("---")
            st.markdown("##### 📥 Exportación Oficial (Formatos Institucionales)")
            col_exp_1, col_exp_2, col_exp_3 = st.columns([2, 1, 1])
            
            est_a_descargar = col_exp_1.selectbox(
                "Seleccionar Estimación a Exportar:", 
                [f"Estimación #{e['num_periodo']} (Del {e['periodo_inicio']} al {e['periodo_fin']})" for e in estimaciones_proyecto]
            )
            formato_institucion = col_exp_2.selectbox("Plantilla Institucional:", ["IMSS", "Poder Judicial de la Federación"])
            proy_obj_actual = next((p for p in lista_proyectos if p["id"] == id_proy_est), {})
            
            with col_exp_3:
                # Div compensatorio para que el botón se alinee con los selectores
                st.markdown("<div style='margin-top: 28px;'></div>", unsafe_allow_html=True)
                btn_generar_excel = st.button("Preparar Archivo .XLSX", type="primary", use_container_width=True)

            if btn_generar_excel:
                idx_est_sel = [f"Estimación #{e['num_periodo']} (Del {e['periodo_inicio']} al {e['periodo_fin']})" for e in estimaciones_proyecto].index(est_a_descargar)
                est_obj_actual = estimaciones_proyecto[idx_est_sel]
                
                if formato_institucion == "IMSS":
                    with st.spinner("Conectando con Supabase e inyectando datos en la plantilla oficial..."):
                        try:
                            # 1. Recuperamos todo el catálogo de este proyecto
                            conceptos_cat = get_conceptos(id_proy_est)
                            
                            # 2. Descargamos la plantilla limpia de Supabase
                            plantilla_bytes = descargar_plantilla_supabase("plantillas", "plantilla_maestra_estimacion_imss.xlsx")
                            
                            # 3. Lanzamos el inyector optimizado
                            xlsx_generado = inyectar_datos_excel_imss(
                                plantilla_bytes=plantilla_bytes,
                                proy_info=proy_obj_actual,
                                estimacion_info=est_obj_actual,
                                conceptos_cat=conceptos_cat,
                                estimaciones_dash=estimaciones_proyecto
                            )
                            
                            # 4. Guardamos el buffer binario localmente en session state
                            st.session_state["xls_buffer"] = xlsx_generado
                            st.session_state["xls_name"] = f"Estimacion_{est_obj_actual['num_periodo']}_{proy_obj_actual.get('contrato_no', 'IMSS')}.xlsx"
                            st.success("✅ Archivo procesado y listo. Haz clic en el botón de abajo para descargarlo.")
                        except Exception as e:
                            st.error(f"Error al generar el formato: {e}")
                else:
                    st.info("🚧 La plantilla para el Poder Judicial de la Federación se encuentra en proceso de homologación.")

            # Si el archivo está en memoria, revelamos el botón nativo de descarga
            if st.session_state.get("xls_buffer"):
                st.download_button(
                    label=f"⬇️ Descargar {st.session_state['xls_name']}",
                    data=st.session_state["xls_buffer"],
                    file_name=st.session_state["xls_name"],
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )

# -------------------------------------------------------------
# TAB 4: CATÁLOGO DEL PROYECTO
# -------------------------------------------------------------
with tab_catalogo:
    st.subheader("Catálogo del Proyecto (Contrato de Obra)")
    if not proyectos_dict: st.warning("Primero debes registrar un proyecto.")
    else:
        proy_sel = st.selectbox("Seleccionar Proyecto Destino", list(proyectos_dict.keys()), key="cat_proy_2")
        proy_id = proyectos_dict[proy_sel]

        categorias_disp = get_biblioteca_categorias()
        cat_sel = st.selectbox("Filtro de Categoría (Para importaciones de Biblioteca):", categorias_disp, key="sel_cat_bib_2")
        conceptos_bib = get_biblioteca_conceptos(cat_sel)

        col_import1, col_import2 = st.columns(2)

        with col_import1:
            with st.expander("📦 Importación Lote desde Biblioteca"):
                if not conceptos_bib: st.info(f"No hay conceptos en '{cat_sel}'.")
                else:
                    opciones_lote = {f"{c['clave']} — {c['descripcion'][:60]}...": c for c in conceptos_bib}
                    seleccionados_lote = st.multiselect("Seleccionar múltiples conceptos:", list(opciones_lote.keys()), key=f"ms_lote_{st.session_state.ms_lote_key}")
                    if st.button("📥 Importar Seleccionados", type="primary") and seleccionados_lote:
                        registros_a_insertar = []
                        for sel in seleccionados_lote:
                            c_ref = opciones_lote[sel]
                            registros_a_insertar.append({
                                "id_proyecto": proy_id, "especialidad": c_ref.get("especialidad", ""),
                                "clave": c_ref["clave"], "descripcion": c_ref["descripcion"], "unidad": normalizar_unidad(c_ref["unidad"]),
                                "cantidad_contratada": 0.0, "precio_unitario": float(c_ref.get("precio_referencial") or 0.0)
                            })
                        if registros_a_insertar:
                            for chunk in [registros_a_insertar[i:i+50] for i in range(0, len(registros_a_insertar), 50)]:
                                supabase.table("catalogo_conceptos").insert(chunk).execute()
                            get_conceptos.clear()
                            st.session_state.ms_lote_key += 1
                            st.success(f"✅ {len(registros_a_insertar)} conceptos importados.")
                            st.rerun()

            with st.expander("📄 Importar Concepto Único desde Biblioteca"):
                if conceptos_bib:
                    opciones_unico = {f"{c['clave']} — {c['descripcion'][:55]}...": c for c in conceptos_bib}
                    obj_u = opciones_unico[st.selectbox("Concepto maestro:", list(opciones_unico.keys()), key="sel_unico")]
                    with st.form("form_importar_bib", clear_on_submit=True):
                        c_cu1, c_cu2 = st.columns(2)
                        cant_u = c_cu1.number_input("Cantidad Contratada", min_value=0.001, value=1.0, step=1.0)
                        pu_u = c_cu2.number_input("Precio Unitario ($)", min_value=0.0, value=float(obj_u.get("precio_referencial") or 0.0), step=10.0)
                        if st.form_submit_button("➕ Agregar al Contrato"):
                            supabase.table("catalogo_conceptos").insert({
                                "id_proyecto": proy_id, "especialidad": obj_u.get("especialidad"),
                                "clave": obj_u["clave"], "descripcion": obj_u["descripcion"], "unidad": normalizar_unidad(obj_u["unidad"]),
                                "cantidad_contratada": cant_u, "precio_unitario": pu_u
                            }).execute()
                            get_conceptos.clear()
                            st.success("Concepto agregado.")
                            st.rerun()

        with col_import2:
            with st.expander("📂 Subir desde Excel (Carga Masiva al Contrato)"):
                st.download_button("📥 Descargar Plantilla Oficial Excel", data=generar_plantilla_excel("proyecto"), file_name="Plantilla_Catalogo_Proyecto.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                excel_proy = st.file_uploader("Sube la plantilla llena:", type=["xlsx"], key=f"up_proy_{st.session_state.up_proy_key}")
                if excel_proy and st.button("Subir e Insertar al Proyecto", type="primary"):
                    try:
                        df_up = pd.read_excel(excel_proy)
                        expected = ["Especialidad", "Clave", "Descripcion", "Unidad", "Cantidad_Contratada", "Precio_Unitario"]
                        if not all(col in df_up.columns for col in expected): st.error("⚠️ Formato incorrecto. Extrae la plantilla base primero.")
                        else:
                            df_up = df_up.fillna("")
                            df_up["Cantidad_Contratada"] = pd.to_numeric(df_up["Cantidad_Contratada"], errors="coerce").fillna(0.0)
                            df_up["Precio_Unitario"] = pd.to_numeric(df_up["Precio_Unitario"], errors="coerce").fillna(0.0)
                            records = []
                            for _, row in df_up.iterrows():
                                records.append({
                                    "id_proyecto": proy_id, "especialidad": str(row["Especialidad"]).strip(),
                                    "clave": str(row["Clave"]).strip(), "descripcion": str(row["Descripcion"]).strip(),
                                    "unidad": normalizar_unidad(str(row["Unidad"])), "cantidad_contratada": float(row["Cantidad_Contratada"]), "precio_unitario": float(row["Precio_Unitario"])
                                })
                            for chunk in [records[i:i+50] for i in range(0, len(records), 50)]:
                                supabase.table("catalogo_conceptos").insert(chunk).execute()
                            get_conceptos.clear()
                            st.session_state.up_proy_key += 1
                            st.success(f"✅ {len(records)} conceptos cargados.")
                            st.rerun()
                    except Exception as e: st.error(f"Error procesando: {e}")

            with st.expander("➕ Alta manual (Concepto Extraordinario)"):
                with st.form("form_concepto_manual", clear_on_submit=True):
                    esp_m = st.text_input("Especialidad", placeholder="Ej: 01 PRELIMINARES")
                    clave_m = st.text_input("Clave de Concepto *", placeholder="Ej: OC01-015-126")
                    unidad_m = st.selectbox("Unidad", unidades_list)
                    desc_m = st.text_area("Descripción detallada *", placeholder="Ingrese la descripción completa...")
                    c_cant_m, c_pu_m = st.columns(2)
                    cant_m = c_cant_m.number_input("Cantidad Contratada", min_value=0.001, value=1.0, step=1.0)
                    pu_m = c_pu_m.number_input("Precio Unitario ($)", min_value=0.0, value=0.0, step=10.0)
                    
                    if st.form_submit_button("Guardar en Catálogo"):
                        if not clave_m.strip() or not desc_m.strip(): st.error("Clave y descripción son obligatorias.")
                        else:
                            supabase.table("catalogo_conceptos").insert({
                                "id_proyecto": proy_id, "especialidad": esp_m.strip(), "clave": clave_m.strip(),
                                "descripcion": desc_m.strip(), "unidad": normalizar_unidad(unidad_m), "cantidad_contratada": cant_m, "precio_unitario": pu_m
                            }).execute()
                            get_conceptos.clear()
                            st.success("✅ Guardado exitosamente.")
                            st.rerun()

        conceptos_proyecto = get_conceptos(proy_id)
        dict_conc_borrar = {f"#{idx} — {c['clave']} ({c['descripcion'][:45]}...)": c for idx, c in enumerate(conceptos_proyecto, start=1)} if conceptos_proyecto else {}

        with st.expander("🗑️ Eliminar Conceptos del Contrato (Borrado en Cascada)"):
            st.warning("Nota: Eliminar conceptos del catálogo borrará irreversiblemente las mediciones y avances físicos asociados a ellos en este proyecto para proteger la conciliación del Dashboard.")
            conc_a_borrar_labels = st.multiselect("Seleccionar conceptos:", list(dict_conc_borrar.keys()), key=f"del_conc_sel_{st.session_state.del_conc_counter}")
            borrar_todos_conc = st.checkbox("⚠️ Selecciona para eliminar todos los conceptos mostrados en la tabla.", key=f"chk_todos_conc_{st.session_state.del_conc_counter}")
            if borrar_todos_conc: conc_a_borrar_labels = list(dict_conc_borrar.keys())
            
            if st.button("Eliminar Seleccionados", type="primary", disabled=len(conc_a_borrar_labels)==0, key="btn_del_cat_bulk"):
                ids_to_delete = [dict_conc_borrar[label]["id"] for label in conc_a_borrar_labels]
                if ids_to_delete:
                    for chunk in [ids_to_delete[i:i+50] for i in range(0, len(ids_to_delete), 50)]:
                        supabase.table("mediciones_campo").delete().in_("id_concepto", chunk).execute()
                        supabase.table("catalogo_conceptos").delete().in_("id", chunk).execute()
                    
                get_conceptos.clear()
                get_mediciones.clear()
                st.session_state.del_conc_counter += 1
                st.success(f"{len(ids_to_delete)} conceptos y sus mediciones han sido eliminados.")
                st.rerun()

        if conceptos_proyecto:
            st.markdown("##### Presupuesto Oficial del Proyecto (Editor Directo):")
            st.caption("💡 *Haz doble clic sobre cualquier celda (clave, descripción, unidad, cantidad o precio) para modificarla directamente.*")
            df_c = pd.DataFrame(conceptos_proyecto)
            df_c.insert(0, "#", range(1, len(df_c) + 1))
            df_c["cantidad_contratada"] = pd.to_numeric(df_c["cantidad_contratada"], errors="coerce").fillna(0.0).astype(float)
            df_c["precio_unitario"] = pd.to_numeric(df_c["precio_unitario"], errors="coerce").fillna(0.0).astype(float)

            edited_catalogo = st.data_editor(
                df_c[["#", "clave", "especialidad", "unidad", "cantidad_contratada", "precio_unitario", "descripcion"]],
                column_config={
                    "#": st.column_config.NumberColumn("#", disabled=True), "clave": st.column_config.TextColumn("Clave"),
                    "especialidad": st.column_config.TextColumn("Especialidad"), "unidad": st.column_config.SelectboxColumn("Unidad", options=unidades_list),
                    "cantidad_contratada": st.column_config.NumberColumn("Cant. Contratada", min_value=0.0, step=1.0),
                    "precio_unitario": st.column_config.NumberColumn("P.U. ($)", format="$%.2f", min_value=0.0, step=10.0),
                    "descripcion": st.column_config.TextColumn("Descripción")
                },
                use_container_width=True, hide_index=True, key="editor_catalogo"
            )

            if "editor_catalogo" in st.session_state and st.session_state["editor_catalogo"].get("edited_rows", {}):
                for row_str, col_vals in st.session_state["editor_catalogo"]["edited_rows"].items():
                    id_conc_mod = conceptos_proyecto[int(row_str)]["id"]
                    up_payload = {}
                    if "clave" in col_vals: up_payload["clave"] = col_vals["clave"].strip()
                    if "especialidad" in col_vals: up_payload["especialidad"] = col_vals["especialidad"].strip()
                    if "unidad" in col_vals: up_payload["unidad"] = normalizar_unidad(col_vals["unidad"])
                    if "cantidad_contratada" in col_vals: up_payload["cantidad_contratada"] = float(col_vals["cantidad_contratada"])
                    if "precio_unitario" in col_vals: up_payload["precio_unitario"] = float(col_vals["precio_unitario"])
                    if "descripcion" in col_vals: up_payload["descripcion"] = col_vals["descripcion"].strip()
                    supabase.table("catalogo_conceptos").update(up_payload).eq("id", id_conc_mod).execute()
                    get_conceptos.clear()
                    st.toast("✅ Concepto actualizado.")
                    st.rerun()

# -------------------------------------------------------------
# TAB 5: BIBLIOTECA MAESTRA GLOBAL
# -------------------------------------------------------------
with tab_biblioteca:
    st.subheader("📖 Biblioteca Maestra de Conceptos (Global)")
    st.caption("Administra el tabulador institucional de precios y claves para usar en múltiples contratos.")
    
    cats_bib = get_biblioteca_categorias()
    
    col_b1, col_b2 = st.columns(2)
    with col_b1:
        with st.container(border=True):
            st.markdown("**1. Selecciona o elimina una categoría existente:**")
            c_bc1, c_bc2 = st.columns([3, 1])
            cat_bib_sel = c_bc1.selectbox("Institución:", cats_bib, key="sel_cat_admin", label_visibility="collapsed")
            c_bc2.button("Eliminar", type="primary", use_container_width=True, on_click=eliminar_categoria_callback)
    with col_b2:
        with st.container(border=True):
            st.markdown("**2. Crea una categoría:**")
            c_nc1, c_nc2 = st.columns([3, 1])
            c_nc1.text_input("Nombre de institución:", placeholder="Ej: ISSSTE, SEDENA...", label_visibility="collapsed", key="input_nueva_cat")
            c_nc2.button("Crear", use_container_width=True, on_click=crear_categoria_callback)

    col_bib_alta, col_bib_del = st.columns(2)
    with col_bib_alta:
        with st.expander("📂 Subir Biblioteca desde Excel (Carga Masiva)"):
            st.download_button("📥 Descargar Plantilla Excel de Biblioteca", data=generar_plantilla_excel("biblioteca"), file_name="Plantilla_Biblioteca_Global.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            excel_bib = st.file_uploader("Sube la plantilla llena:", type=["xlsx"], key=f"up_bib_{st.session_state.up_bib_key}")
            if excel_bib and st.button("Cargar a Biblioteca Global", type="primary"):
                try:
                    df_b = pd.read_excel(excel_bib)
                    if not all(col in df_b.columns for col in ["Especialidad", "Clave", "Descripcion", "Unidad", "Precio_Unitario"]):
                        st.error("⚠️ El formato del Excel no coincide. Descarga la plantilla oficial.")
                    else:
                        df_b = df_b.fillna("")
                        df_b["Precio_Unitario"] = pd.to_numeric(df_b["Precio_Unitario"], errors="coerce").fillna(0.0)
                        records_b = []
                        for _, row in df_b.iterrows():
                            records_b.append({
                                "categoria": cat_bib_sel, "especialidad": str(row["Especialidad"]).strip(),
                                "clave": str(row["Clave"]).strip(), "descripcion": str(row["Descripcion"]).strip(),
                                "unidad": normalizar_unidad(str(row["Unidad"])), "precio_referencial": float(row["Precio_Unitario"])
                            })
                        for chunk in [records_b[i:i+50] for i in range(0, len(records_b), 50)]:
                            supabase.table("biblioteca_conceptos").upsert(chunk, on_conflict="categoria,clave").execute()
                        get_biblioteca_conceptos.clear()
                        get_biblioteca_categorias.clear()
                        st.session_state.up_bib_key += 1
                        st.success(f"✅ {len(records_b)} conceptos guardados en {cat_bib_sel}.")
                        st.rerun()
                except Exception as e: st.error(f"Error procesando: {e}")
        
        with st.expander("➕ Alta manual (Concepto Maestro)"):
            with st.form("form_alta_bib", clear_on_submit=True):
                esp_m = st.text_input("Especialidad", placeholder="Ej: 01 PRELIMINARES")
                clave_m = st.text_input("Clave de Concepto *", placeholder="Ej: OC01-015-126")
                unidad_m = st.selectbox("Unidad", unidades_list)
                desc_m = st.text_area("Descripción detallada *", placeholder="Ingrese la descripción completa...")
                pu_m = st.number_input("Precio Unitario ($)", min_value=0.0, value=0.0, step=10.0)
                if st.form_submit_button("Guardar en Biblioteca"):
                    if not clave_m.strip() or not desc_m.strip(): st.error("Clave y descripción son obligatorias.")
                    else:
                        supabase.table("biblioteca_conceptos").upsert({
                            "categoria": cat_bib_sel, "especialidad": esp_m.strip(), "clave": clave_m.strip(),
                            "descripcion": desc_m.strip(), "unidad": normalizar_unidad(unidad_m), "precio_referencial": pu_m
                        }, on_conflict="categoria,clave").execute()
                        get_biblioteca_conceptos.clear()
                        get_biblioteca_categorias.clear()
                        st.success("✅ Guardado exitosamente.")
                        st.rerun()

    lista_admin_bib = get_biblioteca_conceptos(cat_bib_sel)
    dict_del_bib = {f"#{idx} — {c['clave']} ({c['descripcion'][:50]}...)": c["id"] for idx, c in enumerate(lista_admin_bib, start=1)} if lista_admin_bib else {}

    with col_bib_del:
        with st.expander("🗑️ Eliminar Conceptos Maestros (Borrado Masivo)"):
            sel_del_bib_labels = st.multiselect("Seleccionar conceptos maestros:", list(dict_del_bib.keys()), key=f"del_bib_{st.session_state.bib_del_counter}")
            borrar_toda_cat = st.checkbox("⚠️ Selecciona para eliminar todos los conceptos mostrados en la tabla.", key=f"chk_toda_cat_{st.session_state.bib_del_counter}")
            if borrar_toda_cat: sel_del_bib_labels = list(dict_del_bib.keys())

            if st.button("Eliminar Seleccionados", type="primary", disabled=len(sel_del_bib_labels)==0, key="btn_del_bib_bulk"):
                ids_to_delete = [dict_del_bib[label] for label in sel_del_bib_labels]
                if ids_to_delete:
                    for chunk in [ids_to_delete[i:i+50] for i in range(0, len(ids_to_delete), 50)]:
                        supabase.table("biblioteca_conceptos").delete().in_("id", chunk).execute()
                get_biblioteca_conceptos.clear()
                st.session_state.bib_del_counter += 1
                st.success(f"{len(ids_to_delete)} conceptos eliminados.")
                st.rerun()

    if lista_admin_bib:
        st.markdown(f"##### Conceptos en Categoría: {cat_bib_sel} (Editor Directo):")
        st.caption("💡 *Haz doble clic sobre cualquier celda para modificar la base maestra.*")
        df_admin = pd.DataFrame(lista_admin_bib)
        df_admin.insert(0, "#", range(1, len(df_admin) + 1))
        df_admin["precio_referencial"] = pd.to_numeric(df_admin["precio_referencial"], errors="coerce").fillna(0.0).astype(float)
        
        edited_bib = st.data_editor(
            df_admin[["#", "clave", "especialidad", "unidad", "precio_referencial", "descripcion"]],
            column_config={
                "#": st.column_config.NumberColumn("#", disabled=True), "clave": st.column_config.TextColumn("Clave"),
                "especialidad": st.column_config.TextColumn("Especialidad"), "unidad": st.column_config.SelectboxColumn("Unidad", options=unidades_list),
                "precio_referencial": st.column_config.NumberColumn("Precio Unitario ($)", format="$%.2f", min_value=0.0, step=10.0),
                "descripcion": st.column_config.TextColumn("Descripción")
            },
            use_container_width=True, hide_index=True, key="editor_biblioteca"
        )
        
        if "editor_biblioteca" in st.session_state and st.session_state["editor_biblioteca"].get("edited_rows", {}):
            for row_str, col_vals in st.session_state["editor_biblioteca"]["edited_rows"].items():
                id_bib_mod, up_b = lista_admin_bib[int(row_str)]["id"], {}
                if "clave" in col_vals: up_b["clave"] = col_vals["clave"].strip()
                if "especialidad" in col_vals: up_b["especialidad"] = col_vals["especialidad"].strip()
                if "unidad" in col_vals: up_b["unidad"] = normalizar_unidad(col_vals["unidad"])
                if "precio_referencial" in col_vals: up_b["precio_referencial"] = float(col_vals["precio_referencial"])
                if "descripcion" in col_vals: up_b["descripcion"] = col_vals["descripcion"].strip()
                supabase.table("biblioteca_conceptos").update(up_b).eq("id", id_bib_mod).execute()
                get_biblioteca_conceptos.clear()
                st.toast("✅ Concepto maestro actualizado.")
                st.rerun()

# -------------------------------------------------------------
# TAB 6: PROYECTOS (DATOS GENERALES)
# -------------------------------------------------------------
with tab_proyectos:
    st.subheader("Gestión de Proyectos")
    col_p_alta, col_p_baja = st.columns(2)

    with col_p_alta:
        with st.expander("➕ Dar de alta nuevo proyecto", expanded=False):
            with st.form("form_nuevo_proyecto", clear_on_submit=True):
                nombre_obra = st.text_input("Nombre de la Obra *", placeholder="Inserte aquí el nombre oficial completo de la obra...")
                descripcion_sintetica = st.text_area("Descripción sintética", placeholder="Escriba un resumen del alcance de los trabajos...")
                c1, c2 = st.columns(2)
                ubicacion = c1.text_input("Ubicación", placeholder="Ej: Villa de Álvarez, Colima")
                unidad = c2.text_input("Unidad médica / Inmueble", placeholder="Ej: HGZ-01, UMF-19...")
                c3, c4 = st.columns(2)
                contrato_no = c3.text_input("N° de Contrato", placeholder="Ej: C5M0077")
                concurso_no = c4.text_input("N° de Concurso / Licitación", placeholder="Ej: LO-50-GYR-050GYR080-N-15-2025")
                c5, c6 = st.columns(2)
                contratista = c5.text_input("Contratista / Empresa", placeholder="Razón social o nombre del contratista...")
                residente = c6.text_input("Residente de Obra / Supervisor", placeholder="Nombre del responsable de supervisión...")
                
                if st.form_submit_button("Guardar Proyecto"):
                    if not nombre_obra.strip(): st.error("El nombre de la obra es obligatorio.")
                    else:
                        supabase.table("proyectos").insert({
                            "nombre_obra": nombre_obra.strip(), "descripcion_sintetica": descripcion_sintetica.strip(),
                            "ubicacion": ubicacion.strip(), "unidad": unidad.strip(), "contrato_no": contrato_no.strip(),
                            "concurso_no": concurso_no.strip(), "contratista": contratista.strip(), "residente_obra": residente.strip()
                        }).execute()
                        get_proyectos.clear()
                        st.success("✅ Proyecto registrado con éxito.")
                        st.rerun()

    proy_dict_delete = {f"#{idx} — {p['nombre_obra'][:70]}... ({p.get('contrato_no') or 'S/C'})": p["id"] for idx, p in enumerate(lista_proyectos, start=1)} if lista_proyectos else {}

    with col_p_baja:
        with st.expander("🗑 Eliminar Proyecto"):
            if not proy_dict_delete: st.info("No hay proyectos registrados para eliminar.")
            else:
                proy_del_sel = st.selectbox("Seleccionar proyecto a borrar:", list(proy_dict_delete.keys()), key=f"del_proy_sel_{st.session_state.del_proy_counter}")
                st.warning("⚠️ Eliminar un proyecto borrará en cascada todo su catálogo, estimaciones y mediciones.")
                if st.button("Eliminar Proyecto", type="primary", disabled=not st.checkbox("Confirmo la eliminación definitiva del proyecto", key=f"chk_del_proy_{st.session_state.del_proy_counter}")):
                    supabase.table("proyectos").delete().eq("id", proy_dict_delete[proy_del_sel]).execute()
                    get_proyectos.clear()
                    get_conceptos.clear()
                    get_estimaciones.clear()
                    get_mediciones.clear()
                    st.session_state.del_proy_counter += 1
                    st.success("Proyecto eliminado correctamente.")
                    st.rerun()

    if lista_proyectos:
        st.markdown("##### Proyectos Registrados:")
        st.caption("💡 *Haz doble clic sobre cualquier campo para actualizar los datos oficiales del contrato.*")
        df_p = pd.DataFrame(lista_proyectos)
        df_p.insert(0, "#", range(1, len(df_p) + 1))
        cols_mostrar = ["#", "nombre_obra", "unidad", "contrato_no", "concurso_no", "ubicacion", "contratista", "residente_obra"]
        for c_m in cols_mostrar[1:]: df_p[c_m] = df_p[c_m].fillna("").astype(str)

        edited_proy = st.data_editor(
            df_p[cols_mostrar],
            column_config={
                "#": st.column_config.NumberColumn("#", disabled=True), "nombre_obra": st.column_config.TextColumn("Nombre de la Obra"),
                "unidad": st.column_config.TextColumn("Unidad Médica"), "contrato_no": st.column_config.TextColumn("Contrato N°"),
                "concurso_no": st.column_config.TextColumn("Licitación N°"), "ubicacion": st.column_config.TextColumn("Ubicación"),
                "contratista": st.column_config.TextColumn("Contratista"), "residente_obra": st.column_config.TextColumn("Residente / Supervisor")
            },
            use_container_width=True, hide_index=True, key="editor_proyectos"
        )

        if "editor_proyectos" in st.session_state and st.session_state["editor_proyectos"].get("edited_rows", {}):
            for row_str, col_vals in st.session_state["editor_proyectos"]["edited_rows"].items():
                id_proy_mod, up_p = lista_proyectos[int(row_str)]["id"], {}
                for campo in ["nombre_obra", "unidad", "contrato_no", "concurso_no", "ubicacion", "contratista", "residente_obra"]:
                    if campo in col_vals: up_p[campo] = col_vals[campo].strip()
                supabase.table("proyectos").update(up_p).eq("id", id_proy_mod).execute()
                get_proyectos.clear()
                st.toast("✅ Datos de la obra actualizados.")
                st.rerun()