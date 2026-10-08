import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import io
import datetime
import requests
from modulos.db_engine import normalizar_unidad, get_mediciones

def descargar_plantilla_supabase(bucket_name, file_name):
    supabase = st.session_state["supabase_client"]
    url_plantilla = supabase.storage.from_(bucket_name).get_public_url(file_name)
    response = requests.get(url_plantilla)
    if response.status_code == 200:
        return io.BytesIO(response.content)
    else:
        raise Exception(f"No se pudo descargar la plantilla de {url_plantilla}")

def set_cell_value(ws, col, row, value):
    coord = f"{col}{row}"
    cell = ws[coord]
    if type(cell).__name__ == 'MergedCell':
        for rng in ws.merged_cells.ranges:
            if cell.row >= rng.min_row and cell.row <= rng.max_row and cell.column >= rng.min_col and cell.column <= rng.max_col:
                ws.cell(row=rng.min_row, column=rng.min_col).value = value
                break
    else:
        cell.value = value

def _preparar_datos_estimacion(estimacion_info, conceptos_cat, estimaciones_dash):
    monto_contratado_total = sum(float(c.get("cantidad_contratada") or 0.0) * float(c.get("precio_unitario") or 0.0) for c in conceptos_cat)
    acumulados_anteriores = {c["id"]: 0.0 for c in conceptos_cat}
    cant_periodo_dict = {c["id"]: 0.0 for c in conceptos_cat}
    mediciones_actual = []

    for e in sorted(estimaciones_dash, key=lambda x: x['num_periodo']):
        meds_e = get_mediciones(e["id"])
        if e["id"] == estimacion_info["id"]:
            mediciones_actual = meds_e
            for m in meds_e:
                c_id = m.get("id_concepto")
                cant_periodo_dict[c_id] = cant_periodo_dict.get(c_id, 0.0) + float(m.get("cantidad_total") or 0.0)
        elif e["num_periodo"] < estimacion_info["num_periodo"]:
            for m in meds_e:
                c_id = m.get("id_concepto")
                acumulados_anteriores[c_id] = acumulados_anteriores.get(c_id, 0.0) + float(m.get("cantidad_total") or 0.0)

    conceptos_procesados = []
    for c in conceptos_cat:
        c_id = c["id"]
        cant_ant = acumulados_anteriores.get(c_id, 0.0)
        cant_pres = cant_periodo_dict.get(c_id, 0.0)
        conceptos_procesados.append({
            "id": c_id, 
            "clave": c['clave'], 
            "descripcion": c['descripcion'], 
            "especialidad": c.get('especialidad', '') or 'GENERAL',
            "categoria": c.get('categoria', '') or '',
            "unidad": normalizar_unidad(c['unidad']), 
            "cant_contratada": float(c.get('cantidad_contratada') or 0.0),
            "pu": float(c.get('precio_unitario') or 0.0), 
            "cant_anterior": cant_ant, 
            "cant_periodo": cant_pres, 
            "cant_acumulada": cant_ant + cant_pres
        })
    return monto_contratado_total, conceptos_procesados, mediciones_actual

def inyectar_datos_excel_imss(plantilla_bytes, proy_info, estimacion_info, conceptos_cat, estimaciones_dash):
    wb = openpyxl.load_workbook(plantilla_bytes)
    monto_contratado_total, conceptos_procesados, mediciones_actual = _preparar_datos_estimacion(estimacion_info, conceptos_cat, estimaciones_dash)
    conceptos_con_avance_periodo = [c for c in conceptos_procesados if c['cant_periodo'] > 0]
    
    ws_datos = wb["DATOS"]
    set_cell_value(ws_datos, 'B', 2, proy_info.get('nombre_obra', ''))
    set_cell_value(ws_datos, 'B', 4, proy_info.get('descripcion_sintetica', ''))
    set_cell_value(ws_datos, 'B', 6, proy_info.get('ubicacion', ''))
    set_cell_value(ws_datos, 'B', 8, proy_info.get('unidad', ''))
    set_cell_value(ws_datos, 'B', 12, proy_info.get('contratista', ''))
    set_cell_value(ws_datos, 'B', 14, proy_info.get('residente_obra', ''))
    set_cell_value(ws_datos, 'B', 16, proy_info.get('contrato_no', ''))
    set_cell_value(ws_datos, 'B', 20, f"DEL {estimacion_info.get('periodo_inicio', '')} AL {estimacion_info.get('periodo_fin', '')}")
    set_cell_value(ws_datos, 'B', 23, datetime.datetime.now().strftime("%d/%m/%Y"))
    set_cell_value(ws_datos, 'B', 31, f"{estimacion_info.get('num_periodo', 1)} ({estimacion_info.get('estado', 'NORMAL').upper()})")
    set_cell_value(ws_datos, 'B', 35, monto_contratado_total)

    ws_cat = wb["CATALOGO"]
    fila_cat = 14
    for c in conceptos_cat:
        set_cell_value(ws_cat, 'A', fila_cat, c['clave'])
        set_cell_value(ws_cat, 'B', fila_cat, c['descripcion'])
        set_cell_value(ws_cat, 'C', fila_cat, normalizar_unidad(c['unidad']))
        cant, pu = float(c.get('cantidad_contratada') or 0.0), float(c.get('precio_unitario') or 0.0)
        set_cell_value(ws_cat, 'D', fila_cat, cant)
        set_cell_value(ws_cat, 'E', fila_cat, pu)
        set_cell_value(ws_cat, 'G', fila_cat, cant * pu)
        set_cell_value(ws_cat, 'H', fila_cat, (cant * pu / monto_contratado_total) if monto_contratado_total > 0 else 0)
        fila_cat += 1

    ws_est = wb["Estimacion"]
    fila_est = 14
    for c_idx, c in enumerate(conceptos_con_avance_periodo, start=1):
        set_cell_value(ws_est, 'A', fila_est, c_idx)
        set_cell_value(ws_est, 'B', fila_est, c['descripcion'])
        set_cell_value(ws_est, 'E', fila_est, c['unidad'])
        set_cell_value(ws_est, 'G', fila_est, c['clave'])
        set_cell_value(ws_est, 'H', fila_est, c['cant_contratada'])
        set_cell_value(ws_est, 'I', fila_est, c['cant_acumulada'])
        set_cell_value(ws_est, 'J', fila_est, c['cant_acumulada'] - c['cant_periodo'])
        set_cell_value(ws_est, 'K', fila_est, c['cant_periodo'])
        set_cell_value(ws_est, 'L', fila_est, c['pu'])
        set_cell_value(ws_est, 'M', fila_est, c['cant_periodo'] * c['pu'])
        fila_est += 1

    ws_gen = wb["Generador"]
    fila_gen = 14
    for c in conceptos_con_avance_periodo:
        set_cell_value(ws_gen, 'A', fila_gen, c['clave'])
        set_cell_value(ws_gen, 'B', fila_gen, c['descripcion'])
        set_cell_value(ws_gen, 'J', fila_gen, c['unidad'])
        fila_gen += 1
        
        meds_conc = [m for m in mediciones_actual if m['id_concepto'] == c['id']]
        for m in meds_conc:
            set_cell_value(ws_gen, 'B', fila_gen, m.get('localizacion', ''))
            set_cell_value(ws_gen, 'D', fila_gen, m.get('eje', ''))
            set_cell_value(ws_gen, 'E', fila_gen, m.get('tramo', ''))
            set_cell_value(ws_gen, 'F', fila_gen, float(m.get('ancho') or 0.0))
            set_cell_value(ws_gen, 'G', fila_gen, float(m.get('largo') or 0.0))
            set_cell_value(ws_gen, 'H', fila_gen, float(m.get('alto') or 0.0))
            set_cell_value(ws_gen, 'I', fila_gen, float(m.get('piezas') or 1.0))
            set_cell_value(ws_gen, 'J', fila_gen, c['unidad'])
            set_cell_value(ws_gen, 'K', fila_gen, float(m.get('cantidad_total') or 0.0))
            fila_gen += 1
        fila_gen += 1

    ws_croquis, ws_fotos = wb["Croquis"], wb["Bit. Fotografica"]
    fila_croquis, fila_fotos = 8, 8
    for m in mediciones_actual:
        c_ref = next((c for c in conceptos_cat if c['id'] == m['id_concepto']), {})
        c_clave = c_ref.get('clave', 'S/C')
        if m.get('url_croquis'):
            for u_c in str(m['url_croquis']).split(','):
                u_c = u_c.strip()
                if not u_c: continue
                try:
                    resp = requests.get(u_c, timeout=10)
                    if resp.status_code == 200:
                        img = OpenpyxlImage(io.BytesIO(resp.content))
                        img.width, img.height = 400, 300
                        ws_croquis.add_image(img, f'B{fila_croquis}')
                        set_cell_value(ws_croquis, 'A', fila_croquis, f"Concepto: {c_clave} - Loc: {m.get('localizacion','')}")
                        fila_croquis += 18
                except Exception: pass
        if m.get('url_foto'):
            for u_f in str(m['url_foto']).split(','):
                u_f = u_f.strip()
                if not u_f: continue
                try:
                    resp = requests.get(u_f, timeout=10)
                    if resp.status_code == 200:
                        img = OpenpyxlImage(io.BytesIO(resp.content))
                        img.width, img.height = 400, 300
                        ws_fotos.add_image(img, f'B{fila_fotos}')
                        set_cell_value(ws_fotos, 'A', fila_fotos, f"Concepto: {c_clave} - Loc: {m.get('localizacion','')}")
                        fila_fotos += 18
                except Exception: pass

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def inyectar_datos_excel_pjf(plantilla_bytes, proy_info, estimacion_info, conceptos_cat, estimaciones_dash):
    wb = openpyxl.load_workbook(plantilla_bytes)
    monto_contratado_total, conceptos_procesados, mediciones_actual = _preparar_datos_estimacion(estimacion_info, conceptos_cat, estimaciones_dash)

    ws_car = wb['CARATULA DICIEMBRE 2025']
    set_cell_value(ws_car, 'H', 2, proy_info.get('contrato_no', 'S/N'))
    set_cell_value(ws_car, 'H', 3, datetime.datetime.now().strftime("%d/%m/%Y"))
    set_cell_value(ws_car, 'C', 4, f"ESTIMACIÓN No. {estimacion_info.get('num_periodo', 1):02d}")
    set_cell_value(ws_car, 'C', 5, f"DEL {estimacion_info.get('periodo_inicio', '')} AL {estimacion_info.get('periodo_fin', '')}")
    set_cell_value(ws_car, 'F', 5, proy_info.get('nombre_obra', ''))
    set_cell_value(ws_car, 'F', 8, proy_info.get('ubicacion', ''))
    
    contratista = proy_info.get('contratista', '')
    set_cell_value(ws_car, 'L', 13, contratista if contratista else 'COMPLETAR CONTRATISTA')
    set_cell_value(ws_car, 'D', 22, monto_contratado_total)

    ws_est_cta = wb['ESTADO DE CUENTA AGOSTO 2026']
    telefono = proy_info.get('telefono', '')
    set_cell_value(ws_est_cta, 'E', 6, telefono if telefono else 'COMPLETAR TEL')
    
    fila_cta = 14
    for est in sorted(estimaciones_dash, key=lambda x: x['num_periodo']):
        if est['num_periodo'] > estimacion_info['num_periodo']: continue
        meds = get_mediciones(est['id'])
        importe_est = sum(float(m.get('cantidad_total') or 0.0) * float(next((c['precio_unitario'] for c in conceptos_cat if c['id'] == m.get('id_concepto')), 0)) for m in meds)
        
        set_cell_value(ws_est_cta, 'A', fila_cta, est['num_periodo'])
        set_cell_value(ws_est_cta, 'B', fila_cta, f"Est. {est['num_periodo']} ({est['periodo_inicio']})")
        set_cell_value(ws_est_cta, 'C', fila_cta, importe_est)
        set_cell_value(ws_est_cta, 'I', fila_cta, 0)
        fila_cta += 1
        if fila_cta > 26: break

    ws_res = wb['RESUMEN AGOSTO 2026']
    importes_esp = {}
    for c in conceptos_procesados:
        grupo = c['categoria'] if c['categoria'] else c['especialidad']
        importes_esp[grupo] = importes_esp.get(grupo, 0.0) + (c['cant_periodo'] * c['pu'])
        
    fila_res = 18
    for i, (esp, imp) in enumerate(importes_esp.items(), start=1):
        set_cell_value(ws_res, 'A', fila_res, i)
        set_cell_value(ws_res, 'B', fila_res, esp)
        set_cell_value(ws_res, 'K', fila_res, imp)
        fila_res += 1
        if fila_res > 29: break

    ws_est = wb['ESTIMACIÓN AGOSTO 2026']
    page_starts = [r for r in range(1, ws_est.max_row + 1) if ws_est.cell(row=r, column=4).value and "DATOS DEL CONTRATO" in str(ws_est.cell(row=r, column=4).value).upper()]
    
    c_idx = 0
    slots = [(13, 14), (15, 16), (17, 18), (19, 20)] 
    
    for p_start in page_starts:
        if c_idx >= len(conceptos_procesados): break
        for off_h, off_d in slots:
            if c_idx < len(conceptos_procesados):
                c = conceptos_procesados[c_idx]
                r_hdr = p_start + off_h
                r_data = p_start + off_d
                
                categoria_texto = c['categoria'] if c['categoria'] else c['especialidad']
                set_cell_value(ws_est, 'B', r_hdr, categoria_texto)
                
                set_cell_value(ws_est, 'A', r_data, c['clave'])
                set_cell_value(ws_est, 'B', r_data, c['descripcion'])
                set_cell_value(ws_est, 'E', r_data, c['unidad'])
                set_cell_value(ws_est, 'F', r_data, c['pu'])
                set_cell_value(ws_est, 'G', r_data, c['cant_contratada'])
                set_cell_value(ws_est, 'I', r_data, c['cant_anterior'])
                set_cell_value(ws_est, 'J', r_data, c['cant_periodo'])
                
                c_idx += 1

    ws_fotos = wb['REPORTE FOTOGRAFICO']
    page_starts_fotos = [r for r in range(1, ws_fotos.max_row + 1) if ws_fotos.cell(row=r, column=4).value and "DATOS DEL CONTRATO" in str(ws_fotos.cell(row=r, column=4).value).upper()]
    
    fotos_list = []
    for m in mediciones_actual:
        if m.get('url_foto') or m.get('url_croquis'):
            c_ref = next((c for c in conceptos_cat if c['id'] == m['id_concepto']), {})
            desc_base = f"{c_ref.get('clave', '')} - {m.get('localizacion', '')}"
            if m.get('url_foto'):
                for u_f in str(m['url_foto']).split(','):
                    if u_f.strip(): fotos_list.append((u_f.strip(), desc_base))
            if m.get('url_croquis'):
                for u_c in str(m['url_croquis']).split(','):
                    if u_c.strip(): fotos_list.append((u_c.strip(), f"{desc_base} (Croquis)"))

    def _insertar_img(ws, r_img, c_let, r_txt, f_data):
        url, desc = f_data
        try:
            resp = requests.get(url, timeout=10)
            if resp.status_code == 200:
                img = OpenpyxlImage(io.BytesIO(resp.content))
                img.width, img.height = 320, 240
                ws.add_image(img, f"{c_let}{r_img}")
                ws.cell(row=r_txt, column=1 if c_let == 'B' else 8, value=desc)
        except Exception: pass

    f_idx = 0
    for p_start in page_starts_fotos:
        if f_idx >= len(fotos_list): break
        slots_img = [('B', 16, 15), ('H', 16, 15), ('B', 25, 24), ('H', 25, 24)]
        for c_let, r_img, r_txt in slots_img:
            if f_idx < len(fotos_list):
                _insertar_img(ws_fotos, p_start + r_img, c_let, p_start + r_txt, fotos_list[f_idx])
                f_idx += 1

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()

def generar_excel_estimapp(proy_info, estimacion_info, conceptos_cat, estimaciones_dash):
    monto_contratado_total, conceptos_procesados, mediciones_actual = _preparar_datos_estimacion(estimacion_info, conceptos_cat, estimaciones_dash)
    
    try:
        dt = pd.to_datetime(estimacion_info.get('periodo_fin'))
        meses_corto = ["ENE", "FEB", "MAR", "ABR", "MAY", "JUN", "JUL", "AGO", "SEP", "OCT", "NOV", "DIC"]
        mes_anio_corto = f"{meses_corto[dt.month - 1]} {str(dt.year)[-2:]}"
    except:
        mes_anio_corto = "ACTUAL"

    wb = openpyxl.Workbook()
    font_titulo = Font(name="Segoe UI", size=14, bold=True, color="1F2937")
    font_headers = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_datos = Font(name="Segoe UI", size=9, color="111827")
    font_bold = Font(name="Segoe UI", size=9, bold=True, color="111827")
    
    fill_primario = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_secundario = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    fill_totales = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")
    
    borde_delgado = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000")
    )

    def draw_header(ws, tab_name):
        ws.views.sheetView[0].showGridLines = False
        ws.merge_cells("A1:K1")
        ws["A1"] = f"🏗️ ESTIMAPP | {tab_name} | Estimación #{estimacion_info.get('num_periodo', 1):02d}"
        ws["A1"].font = font_titulo
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws["A1"].border = borde_delgado
        ws.row_dimensions[1].height = 30
        
        ws["A3"], ws["B3"] = "Proyecto / Obra:", proy_info.get("nombre_obra", "")
        ws["A4"], ws["B4"] = "No. Contrato:", proy_info.get("contrato_no", "")
        ws["A5"], ws["B5"] = "Ubicación:", proy_info.get("ubicacion", "")
        ws["H3"], ws["I3"] = "Contratista:", proy_info.get("contratista", "")
        ws["H4"], ws["I4"] = "Periodo:", f"{estimacion_info.get('periodo_inicio', '')} al {estimacion_info.get('periodo_fin', '')}"
        ws["H5"], ws["I5"] = "Monto Contratado:", monto_contratado_total
        ws["I5"].number_format = '"$"#,##0.00'
        
        for r in [3, 4, 5]:
            ws[f"A{r}"].font = ws[f"H{r}"].font = font_bold
            ws[f"B{r}"].font = ws[f"I{r}"].font = font_datos

    ws_res = wb.active
    ws_res.title = f"Resumen {mes_anio_corto}".strip()
    draw_header(ws_res, "RESUMEN FINANCIERO")
    
    subtotal_est = sum(c['cant_periodo'] * c['pu'] for c in conceptos_procesados)
    iva = subtotal_est * 0.16
    total_liquido = subtotal_est + iva

    bloque_finanzas = [
        ("Subtotal de los Trabajos Ejecutados:", subtotal_est),
        ("I.V.A. (16%):", iva),
        ("Total Líquido a Pagar:", total_liquido)
    ]
    
    fila_f = 8
    for label, val in bloque_finanzas:
        ws_res.merge_cells(start_row=fila_f, start_column=3, end_row=fila_f, end_column=6)
        ws_res.cell(row=fila_f, column=3, value=label).font = font_bold
        ws_res.cell(row=fila_f, column=3).alignment = Alignment(horizontal="right")
        c_val = ws_res.cell(row=fila_f, column=7, value=val)
        c_val.font = font_datos
        c_val.number_format = '"$"#,##0.00'
        c_val.border = borde_delgado
        if "Líquido" in label:
            c_val.fill = fill_totales
            c_val.font = font_bold
        fila_f += 2

    ws_cta = wb.create_sheet(f"Est Cta {mes_anio_corto}".strip())
    draw_header(ws_cta, "ESTADO DE CUENTA HISTÓRICO")
    
    encabezados_cta = ["No. Est.", "Periodo de Estimación", "Subtotal Ejecutado", "I.V.A.", "Total Líquido"]
    for col_idx, col_nombre in enumerate(encabezados_cta, start=1):
        c = ws_cta.cell(row=7, column=col_idx, value=col_nombre)
        c.font, c.fill, c.alignment, c.border = font_headers, fill_primario, Alignment(horizontal="center"), borde_delgado
    
    fila_c = 8
    for est in sorted(estimaciones_dash, key=lambda x: x['num_periodo']):
        if est['num_periodo'] > estimacion_info['num_periodo']: continue
        meds = get_mediciones(est['id'])
        subt = sum(float(m.get('cantidad_total') or 0.0) * float(next((c['precio_unitario'] for c in conceptos_cat if c['id'] == m.get('id_concepto')), 0)) for m in meds)
        imp_iva = subt * 0.16
        liquido = subt + imp_iva
        
        ws_cta.cell(row=fila_c, column=1, value=est['num_periodo']).alignment = Alignment(horizontal="center")
        ws_cta.cell(row=fila_c, column=2, value=f"{est['periodo_inicio']} al {est['periodo_fin']}")
        ws_cta.cell(row=fila_c, column=3, value=subt).number_format = '"$"#,##0.00'
        ws_cta.cell(row=fila_c, column=4, value=imp_iva).number_format = '"$"#,##0.00'
        ws_cta.cell(row=fila_c, column=5, value=liquido).number_format = '"$"#,##0.00'
        
        for col_idx in range(1, 6):
            ws_cta.cell(row=fila_c, column=col_idx).font = font_datos
            ws_cta.cell(row=fila_c, column=col_idx).border = borde_delgado
        fila_c += 1

    ws_gen = wb.create_sheet(f"Estimación {mes_anio_corto}".strip())
    draw_header(ws_gen, "CUERPO DE ESTIMACIÓN Y AVANCES")
    
    columnas = ["Clave", "Categoría", "Descripción de Concepto", "Unidad", "P.U.", "Vol. Contrato", "Vol. Acum. Ant.", "Vol. Estimado", "Vol. Acum. Act.", "Vol. Faltante", "Importe Estimado", "% Avance"]
    for col_idx, col_nombre in enumerate(columnas, start=1):
        c = ws_gen.cell(row=7, column=col_idx, value=col_nombre)
        c.font, c.fill, c.alignment, c.border = font_headers, fill_primario, Alignment(horizontal="center", vertical="center", wrap_text=True), borde_delgado
        
    fila_act = 8
    for item in conceptos_procesados:
        if item['cant_periodo'] <= 0 and item['cant_contratada'] == 0: continue
        pu, vol_contrato = item['pu'], item['cant_contratada']
        vol_ant, vol_est, vol_act = item['cant_anterior'], item['cant_periodo'], item['cant_acumulada']
        
        ws_gen.cell(row=fila_act, column=1, value=item['clave']).alignment = Alignment(horizontal="center")
        ws_gen.cell(row=fila_act, column=2, value=item['categoria']).alignment = Alignment(horizontal="center")
        ws_gen.cell(row=fila_act, column=3, value=item['descripcion']).alignment = Alignment(wrap_text=True)
        ws_gen.cell(row=fila_act, column=4, value=item['unidad']).alignment = Alignment(horizontal="center")
        ws_gen.cell(row=fila_act, column=5, value=pu).number_format = '"$"#,##0.00'
        ws_gen.cell(row=fila_act, column=6, value=vol_contrato).number_format = '#,##0.00'
        ws_gen.cell(row=fila_act, column=7, value=vol_ant).number_format = '#,##0.00'
        ws_gen.cell(row=fila_act, column=8, value=vol_est).number_format = '#,##0.00'
        ws_gen.cell(row=fila_act, column=9, value=vol_act).number_format = '#,##0.00'
        ws_gen.cell(row=fila_act, column=10, value=vol_contrato - vol_act).number_format = '#,##0.00'
        ws_gen.cell(row=fila_act, column=11, value=vol_est * pu).number_format = '"$"#,##0.00'
        ws_gen.cell(row=fila_act, column=12, value=(vol_act/vol_contrato) if vol_contrato > 0 else 0).number_format = '0.00%'
        
        for col_idx in range(1, 13):
            cell = ws_gen.cell(row=fila_act, column=col_idx)
            cell.font, cell.border = font_datos, borde_delgado
            if fila_act % 2 == 0: cell.fill = fill_secundario
        fila_act += 1
        
    ws_fotos = wb.create_sheet(f"Fotos {mes_anio_corto}".strip())
    draw_header(ws_fotos, "EVIDENCIA FOTOGRÁFICA")
    
    fila_foto = 8
    col_foto = 2
    for m in mediciones_actual:
        if m.get('url_foto') or m.get('url_croquis'):
            c_ref = next((c for c in conceptos_cat if c['id'] == m['id_concepto']), {})
            desc_base = f"{c_ref.get('clave', '')} - {m.get('localizacion', '')}"
            urls = []
            if m.get('url_foto'):
                for u_f in str(m['url_foto']).split(','):
                    if u_f.strip(): urls.append((u_f.strip(), desc_base))
            if m.get('url_croquis'):
                for u_c in str(m['url_croquis']).split(','):
                    if u_c.strip(): urls.append((u_c.strip(), f"{desc_base} (Croquis)"))
            
            for url, desc in urls:
                try:
                    resp = requests.get(url, timeout=10)
                    if resp.status_code == 200:
                        img = OpenpyxlImage(io.BytesIO(resp.content))
                        img.width, img.height = 320, 240
                        ws_fotos.add_image(img, f"{get_column_letter(col_foto)}{fila_foto}")
                        
                        celda_texto = ws_fotos.cell(row=fila_foto + 13, column=col_foto, value=desc)
                        celda_texto.font = font_bold
                        celda_texto.border = borde_delgado
                        
                        col_foto += 5
                        if col_foto > 7:
                            col_foto = 2
                            fila_foto += 16
                except: pass

    for ws_ajuste in [ws_res, ws_cta, ws_gen, ws_fotos]:
        ws_ajuste.column_dimensions['A'].width = 14
        ws_ajuste.column_dimensions['B'].width = 25
        ws_ajuste.column_dimensions['C'].width = 40
        for col_letter in ['D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L']: 
            ws_ajuste.column_dimensions[col_letter].width = 15

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()